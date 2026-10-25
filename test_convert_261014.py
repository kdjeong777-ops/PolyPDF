# -*- coding: utf-8 -*-
"""261014-11: 형식 변환 — Excel(표)·텍스트·이미지와 공통 바탕 (형식 변환 SOT §1~§4, 개선 로드맵 P2-1·P2-2).

실제 처리기(`action_convert` — 메뉴 넷이 모두 들어오는 곳)로:
A. Excel — 표마다 시트, 칸 글이 그대로
B. 텍스트 — 텍스트 창과 같은 원천(고침 반영), UTF-8 BOM, 쪽 사이 빈 줄 둘
C. 이미지 — 같은 이름의 폴더에 쪽마다 PNG
D. 같은 이름이 있으면 ` (2)` 부터
E. 쪽 범위 `2-3`·썸네일에서 고른 쪽만
F. 취소하면 만들던 파일을 지운다 · 표가 없으면 Excel 을 만들지 않는다
G. 변환은 메인이 아닌 스레드에서 돈다
H. 일괄(책갈피창 여러 파일) — 파일마다 결과
I. 메뉴 넷 — 파일 '내보내기', 도구 막대 '변환', 책갈피창 우클릭 '변환'(261016-1 → 261025-5), 썸네일 우클릭 '선택한 쪽 변환'
"""
import os, sys, tempfile, shutil, threading
os.environ["QT_QPA_PLATFORM"] = "offscreen"
sys.stdout.reconfigure(encoding="utf-8")
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from pathlib import Path

import fitz
from PyQt6.QtCore import QStandardPaths
QStandardPaths.setTestModeEnabled(True)
from PyQt6.QtWidgets import QApplication, QMessageBox, QMenu
from PyQt6.QtTest import QTest

app = QApplication.instance() or QApplication(sys.argv[:1])
fails = []


def chk(cond, msg, extra=""):
    print(("PASS" if cond else "FAIL"), "-", msg, extra if not cond else "")
    if not cond:
        fails.append(msg)


def spin(n=10, ms=20):
    for _ in range(n):
        app.processEvents(); QTest.qWait(ms)


def make_pdf(path, tag, table=True, pages=3):
    d = fitz.open()
    for i in range(pages):
        pg = d.new_page(width=595, height=842)
        pg.insert_text((72, 80), f"{tag} heading {i + 1}", fontsize=16)
        pg.insert_text((72, 110), f"{tag} body line of page {i + 1}", fontsize=11)
        if table and i == 0:
            x0, y0, cw, rh = 72, 160, 140, 24
            for r in range(4):
                pg.draw_line((x0, y0 + r * rh), (x0 + 3 * cw, y0 + r * rh))
            for c in range(4):
                pg.draw_line((x0 + c * cw, y0), (x0 + c * cw, y0 + 3 * rh))
            for r in range(3):
                for c in range(3):
                    pg.insert_text((x0 + c * cw + 6, y0 + r * rh + 16), f"R{r}C{c}", fontsize=10)
    d.save(str(path)); d.close()


root = Path(tempfile.mkdtemp(prefix="polypdf_conv_"))
try:
    a = root / "A.pdf"; b = root / "B.pdf"; plain = root / "P.pdf"
    make_pdf(a, "AAA"); make_pdf(b, "BBB"); make_pdf(plain, "PPP", table=False)
    QMessageBox.question = staticmethod(lambda *x, **k: QMessageBox.StandardButton.No)
    infos = []
    QMessageBox.information = staticmethod(lambda *x, **k: infos.append(x[2] if len(x) > 2 else ""))
    QMessageBox.warning = staticmethod(lambda *x, **k: infos.append(x[2] if len(x) > 2 else ""))
    from viewer.widgets import convert_dialog as cdmod
    setup = {"fn": None}

    def fake_exec(dlg):
        if setup["fn"]:
            setup["fn"](dlg)
        dlg._accept()
        return dlg.result() or cdmod.ConvertDialog.DialogCode.Accepted
    cdmod.ConvertDialog.exec = fake_exec

    from viewer.app import MainWindow
    from viewer.history import HistoryItem
    mw = MainWindow(); mw.resize(1300, 850); mw.show(); app.processEvents()
    mw.open_folder(root); spin(15)
    mw._load_main(HistoryItem(str(a), 0, "", "bookmark")); spin(10)

    # ── G 를 함께 본다: 쓰기 함수가 어느 스레드에서 도는가 ─────────────────
    from viewer.convert import writers
    seen_thread = []
    for k in ("xlsx", "txt", "png"):
        real = writers.WRITERS[k]

        def spy(*aa, _r=real, **kk):
            seen_thread.append(threading.current_thread() is threading.main_thread())
            return _r(*aa, **kk)
        writers.WRITERS[k] = spy

    # ── A. Excel ──────────────────────────────────────────────────
    ok = mw.action_convert(fmt="xlsx"); spin(3)
    x = root / "A.xlsx"
    chk(ok and x.exists(), "A Excel 파일이 같은 폴더·같은 이름으로")
    if x.exists():
        from openpyxl import load_workbook
        wb = load_workbook(str(x))
        vals = [[c.value for c in r] for r in wb[wb.sheetnames[0]].iter_rows()]
        chk(wb.sheetnames == ["p1_1"] and vals and vals[0][:3] == ["R0C0", "R0C1", "R0C2"]
            and vals[2][2] == "R2C2", "A 표마다 시트, 칸 글이 그대로", f"{wb.sheetnames} {vals[:3]}")

    # ── B. 텍스트 + 고침 ──────────────────────────────────────────
    from viewer import text_extract2 as tx
    from viewer.text_fix_store import store
    d = fitz.open(str(a)); rows = tx.page_lines(d, str(a), 1, tables="lines"); d.close()
    t = next((i, r) for i, r in enumerate(rows) if "body line of page 2" in r.get("text", ""))
    store().set_fix(str(a), 1, t[0], "FIXED BODY TWO", orig=t[1]["text"], rect=t[1].get("rect"))
    mw.action_convert(fmt="txt"); spin(3)
    tf = root / "A.txt"
    raw = tf.read_bytes() if tf.exists() else b""
    txt = raw.decode("utf-8-sig").replace(chr(13) + chr(10), chr(10)) if raw else ""   # Windows 줄 끝(CRLF) - 메모장 친화
    chk(raw[:3] == b"\xef\xbb\xbf", "B 텍스트는 UTF-8 BOM")
    chk("AAA heading 3" in txt and "FIXED BODY TWO" in txt and "body line of page 2" not in txt,
        "B 텍스트 창과 같은 원천 — 고침이 반영된다", txt[:120])
    chk(txt.count("\n\n\n") >= 2, "B 쪽 사이 빈 줄 둘")

    # ── C. 이미지 ─────────────────────────────────────────────────
    mw.action_convert(fmt="png"); spin(3)
    imgdir = root / "A"
    pngs = sorted(p.name for p in imgdir.glob("*.png")) if imgdir.is_dir() else []
    chk(pngs == ["A_p0001.png", "A_p0002.png", "A_p0003.png"], "C 같은 이름의 폴더에 쪽마다 PNG", str(pngs))

    # ── D. 이름 겹침 ──────────────────────────────────────────────
    mw.action_convert(fmt="txt"); spin(3)
    chk((root / "A (2).txt").exists(), "D 같은 이름이 있으면 ' (2)' 부터")

    # ── E. 쪽 범위 · 고른 쪽 ───────────────────────────────────────
    def pick_range(dlg):
        dlg.rb_range.setChecked(True); dlg.ed_range.setText("2-3")
    setup["fn"] = pick_range
    mw.action_convert(fmt="txt"); spin(3)
    setup["fn"] = None
    t3 = (root / "A (3).txt").read_text(encoding="utf-8-sig") if (root / "A (3).txt").exists() else ""
    chk("AAA heading 2" in t3 and "AAA heading 3" in t3 and "AAA heading 1" not in t3, "E 쪽 범위 2-3 만", t3[:80])
    mw.action_convert(fmt="txt", pages=[0]); spin(3)                       # 썸네일에서 고른 쪽
    t4 = (root / "A (4).txt").read_text(encoding="utf-8-sig") if (root / "A (4).txt").exists() else ""
    chk("AAA heading 1" in t4 and "AAA heading 2" not in t4, "E 썸네일에서 고른 쪽만", t4[:80])

    # ── F. 취소 · 표 없음 ─────────────────────────────────────────
    from viewer import convert
    n_thread = len(seen_thread)                                        # 아래 직접 호출(메인)은 G 에서 뺀다
    try:
        convert.run(str(a), "png", progress=lambda i, n, l: i < 1)       # 둘째 쪽에서 취소
        cancelled = False
    except convert.Cancelled:
        cancelled = True
    chk(cancelled and not (root / "A (2)").exists(), "F 취소하면 만들던 결과(이미지 폴더)를 지운다")
    infos.clear()
    mw._load_main(HistoryItem(str(plain), 0, "", "bookmark")); spin(8)
    mw.action_convert(fmt="xlsx"); spin(3)
    chk(not (root / "P.xlsx").exists() and any("표가 없습니다" in str(m) for m in infos),
        "F 표가 없으면 Excel 을 만들지 않고 알린다", str(infos[-1:]))

    # ── G ────────────────────────────────────────────────────────
    app_calls = seen_thread[:n_thread] + seen_thread[n_thread + 1:]
    chk(app_calls and not any(app_calls), "G 앱의 변환은 메인이 아닌 스레드에서 돈다", str(seen_thread))

    # ── H. 일괄 ──────────────────────────────────────────────────
    mw.bookmark_tree.convertFilesRequested.emit([str(a), str(b)], "txt"); spin(5)
    chk((root / "B.txt").exists() and (root / "A (5).txt").exists(), "H 일괄 — 파일마다 결과")

    # ── I. 메뉴 넷 ────────────────────────────────────────────────
    want_fmt = [f.label for f in convert.ready_formats()]          # 다 만든 형식이 늘면 함께 는다
    # 261017-12(형식 변환 SOT §2.1, 사용자 지시): Word 셋은 붙어 있고 'Word (Microsoft Word 사용)' 은 '모양 그대로' 바로 아래
    _keys = [f.key for f in convert.FORMATS]
    chk(_keys[:3] == ["docx_edit", "docx_shape", "docx_word"], "I 형식 차례 — Word 셋이 붙어 있다", str(_keys[:4]))
    chk(mw.act_export.text() == "변환..." and mw.act_export.menu() is None,
        "I 파일 메뉴 '변환...' — 하위 메뉴 없이 창(261014-28, 261025-5 이름)")
    chk(mw.act_export.iconText() == "변환", "I 모음·파일·도구 탭 '변환' 단추(261024-1 — 메뉴 액션의 짧은 이름, 261025-5)")
    menus = []
    _orig = QMenu.exec

    def cap(self, *aa, **kk):
        menus.append([(x.text(), [y.text() for y in x.menu().actions()] if x.menu() else None) for x in self.actions()])
        return None
    QMenu.exec = cap
    try:
        bt = mw.bookmark_tree
        node = next(n for n in bt._iter_file_nodes() if Path(n.data(0, bt.DATA_FILE)).name == "A.pdf")
        bt._on_tree_context_menu(bt.tree.visualItemRect(node).center()); spin(2)
        mw._load_main(HistoryItem(str(a), 0, "", "bookmark")); spin(8)
        L = mw.page_thumbs.list
        L.clearSelection(); L.item(0).setSelected(True)
        mw.page_thumbs._on_list_menu(L.visualItemRect(L.item(0)).center()); spin(2)
    finally:
        QMenu.exec = _orig
    # 261016-1(사용자 지시): 책갈피창은 하위 메뉴 없이 '변환...'(261025-5 이름) — 형식은 변환 창에서
    tree_it = [s for t, s in (menus[0] if menus else []) if t == "변환..."]
    chk(tree_it == [None], "I 책갈피창 우클릭 '변환...'(하위 메뉴 없음)", str(menus[:1]))
    th_sub = [s for t, s in (menus[1] if len(menus) > 1 else []) if t.startswith("선택한 쪽 변환")]
    chk(th_sub and th_sub[0] == want_fmt, "I 썸네일 우클릭 '선택한 쪽 변환 ▸'",
        str(menus[1:2]))
finally:
    try:
        from viewer.text_fix_store import store as _s
        _s()._data = {}; _s().save()
    except Exception:
        pass
    shutil.rmtree(root, ignore_errors=True)

print("\n=== " + ("ALL PASS" if not fails else f"FAILURE ({len(fails)})") + " ===")
for m in fails:
    print(" -", m)
sys.stdout.flush()
os._exit(1 if fails else 0)
