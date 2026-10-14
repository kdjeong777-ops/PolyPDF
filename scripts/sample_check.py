# -*- coding: utf-8 -*-
"""표본 시험 — 내보내기(형식 변환)·압축·가림·비교를 실제 업무 PDF 여러 종으로 (릴리스 SOT §4.4, 261014-29).

    python scripts\\sample_check.py [--suite ..\\sample_check_suite.json] [--only S1,S4] [--formats docx_edit,xlsx] [--no-com]

- 표본·쪽은 `sample_check_suite.json`(업무 파일 이름 — 공개 저장소 밖). 표본은 **작업 폴더에 복사**해 쓴다(원본 옆에 아무것도 만들지 않는다).
- 설정 폴더는 시험 모드(`QStandardPaths.setTestModeEnabled`) — 사용자 OCR·고침·색인을 읽거나 쓰지 않는다.
- 형식마다: 앱과 같은 `convert.run`(배경 스레드 대신 이 프로세스) → **라이브러리로 열기**(쪽·슬라이드·시트·글 수) →
  `--no-com` 이 아니면 **Word·한글·PowerPoint 로 열어 PDF 로 찍고**(COM) 원본 쪽 낱말이 결과에 얼마나 남았나(글 재현율 — 차례·띄어쓰기와 무관)와 쪽 수를 본다.
- 압축·가림·비교는 작은 표본(쪽 100 이하)에서: 압축은 결과가 열리고 쪽 수·글이 같은지, 가림은 찾은 낱말이 결과에서 뽑히지 않는지,
  비교는 한 쪽에 글을 더한 사본에서 그 쪽의 '넣음' 을 찾는지.
- 결과: `..\\_review\\samples_<버전>_<시각>\\summary.md`(+ `results.json`, 만든 파일은 `work\\`). 판정은 종료 코드(실패 있으면 1).
"""
from __future__ import annotations

import json
import os
import re
import shutil
import subprocess
import sys
import time
import traceback
import zipfile
from pathlib import Path

HERE = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(HERE))
os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

RECALL_MIN = {  # 원본 쪽 낱말이 결과 글에 있는 비율의 하한 —
              "docx_edit": 0.85, "docx_shape": 0.85, "hwpx_edit": 0.85, "hwpx_shape": 0.85, "pptx_shape": 0.85,
              "docx_word": 0.80, "txt": 0.90}           # 글자층이 있는 쪽만
COM_TIMEOUT_S = 240
INFO_ONLY = ("docx_word",)     # Word 가 직접 바꾸는 형식 — 글 재현은 Word 의 몫이라 참고로만(글이 아예 없으면 앱이 이미 실패로 알린다)
SMALL_PAGES = 100                                     # 압축·가림·비교를 돌리는 표본 크기 상한


_KEEP = re.compile(r"[0-9a-z가-힣]")
_SAME = str.maketrans({"o": "0", "l": "1", "i": "1"})       # OCR 글자층의 O↔0·l↔1 — 양쪽에 같게(우리 정리가 고친 것을 '빠짐' 으로 세지 않게)


def _stream(text: str) -> str:
    return "".join(_KEEP.findall(str(text or "").lower())).translate(_SAME)


def words(text: str) -> set:
    """원본 글의 낱말(2글자 이상, 소문자·O↔0 같게)."""
    return {w for w in (_stream(t) for t in str(text or "").split()) if len(w) >= 2}


def recall(src_words: set, out_text: str) -> float:
    """원본 낱말 중 결과 글(빈칸을 뺀 한 줄)에 들어 있는 비율 — 차례·띄어쓰기·숫자 붙이기와 무관하다.
    표 칸을 행·열 어느 차례로 읽든, OCR 글자층의 `1 5998` 을 `15998` 로 이어도 같게 센다."""
    if not src_words:
        return 1.0
    st = _stream(out_text)
    return sum(1 for w in src_words if w in st) / len(src_words)


# ── 결과 파일을 라이브러리로 열기 ─────────────────────────────
def lib_read(key: str, path: Path) -> dict:
    if key.startswith("docx"):
        import docx
        d = docx.Document(str(path))
        t = "\n".join(p.text for p in d.paragraphs)
        for tb in d.tables:
            for row in tb.rows:
                t += "\n" + " ".join(c.text for c in row.cells)
        for sec in d.sections:
            t += "\n" + " ".join(p.text for p in sec.header.paragraphs) + " " + " ".join(p.text for p in sec.footer.paragraphs)
        return {"paras": len(d.paragraphs), "tables": len(d.tables), "images": len(d.inline_shapes), "text": t}
    if key.startswith("hwpx"):
        with zipfile.ZipFile(path) as z:
            secs = [n for n in z.namelist() if re.match(r"Contents/section\d+\.xml", n)]
            xml = "".join(z.read(n).decode("utf-8", "replace") for n in secs)
            imgs = [n for n in z.namelist() if n.startswith("BinData/")]
        t = " ".join(re.findall(r"<hp:t>([^<]*)</hp:t>", xml))
        return {"sections": len(secs), "tables": xml.count("<hp:tbl "), "images": len(imgs), "text": t}
    if key == "pptx_shape":
        import pptx
        p = pptx.Presentation(str(path))
        t = []
        for s in p.slides:
            for sh in s.shapes:
                if sh.has_text_frame:
                    t.append(sh.text_frame.text)
                if getattr(sh, "has_table", False) and sh.has_table:
                    t += [c.text for r in sh.table.rows for c in r.cells]
        return {"slides": len(p.slides), "text": "\n".join(t)}
    if key == "xlsx":
        import openpyxl
        wb = openpyxl.load_workbook(str(path), read_only=True)
        t = []
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                t += [str(v) for v in row if v is not None]
        return {"sheets": len(wb.worksheets), "text": " ".join(t)}
    if key == "txt":
        return {"text": path.read_text(encoding="utf-8", errors="replace")}
    if key == "png":
        files = sorted(path.glob("*.png"))
        return {"files": len(files), "text": ""}
    return {"text": ""}


# ── COM 으로 열어 PDF 로 찍기 (한 프로그램을 자식 파이썬에서 — 멈추면 시간 상한으로 끝낸다) ──
COM_SCRIPT = r'''
import sys, json
import win32com.client as w
kind, src, dst = sys.argv[1], sys.argv[2], sys.argv[3]
out = {}
if kind == "word":
    app = w.DispatchEx("Word.Application"); app.Visible = False; app.DisplayAlerts = 0
    try:
        d = app.Documents.Open(src, ReadOnly=True, AddToRecentFiles=False)
        d.ExportAsFixedFormat(dst, 17); out["pages"] = d.ComputeStatistics(2); d.Close(False)
    finally:
        app.Quit()
elif kind == "hwp":
    # 한글 하나로 모든 파일을(인자 = 입력 출력 쌍 여럿) — 한글은 파일마다 '접근 허용' 보안 창을 띄우는데, 한 번 띄운
    # 한글에서는 '모두 허용' 한 번이면 끝난다(파일마다 새 한글이면 매번 묻는다, 261014-30)
    hwp = w.Dispatch("HWPFrame.HwpObject")
    try:
        hwp.RegisterModule("FilePathCheckDLL", "FilePathCheckerModule")
    except Exception:
        pass
    try:
        pairs = sys.argv[2:]
        for k, (s_, d_) in enumerate(zip(pairs[::2], pairs[1::2])):
            o = {"i": k}
            try:
                o["open"] = bool(hwp.Open(s_, "HWPX", "forceopen:true")); o["pages"] = hwp.PageCount
                o["save"] = bool(hwp.SaveAs(d_, "PDF", "")); hwp.Clear(1)
            except Exception as e:
                o["err"] = repr(e)[:300]
            print("RESULT " + json.dumps(o), flush=True)
    finally:
        hwp.Quit()
    sys.exit(0)
elif kind == "ppt":
    app = w.DispatchEx("PowerPoint.Application")
    try:
        p = app.Presentations.Open(src, ReadOnly=True, Untitled=False, WithWindow=False)
        p.SaveAs(dst, 32); out["pages"] = p.Slides.Count; p.Close()
    finally:
        app.Quit()
print("RESULT " + json.dumps(out))
'''


def com_installed(kind: str) -> bool:
    import winreg
    key = {"word": r"Word.Application\CLSID", "hwp": r"HWPFrame.HwpObject\CLSID", "ppt": r"PowerPoint.Application\CLSID"}[kind]
    try:
        winreg.CloseKey(winreg.OpenKey(winreg.HKEY_CLASSES_ROOT, key))
        return True
    except OSError:
        return False


def com_to_pdf(kind: str, src: Path, dst: Path, script: Path) -> dict:
    t0 = time.monotonic()
    try:
        r = subprocess.run([sys.executable, str(script), kind, str(src), str(dst)], capture_output=True, text=True,
                           encoding="utf-8", errors="replace", timeout=COM_TIMEOUT_S)
    except subprocess.TimeoutExpired:
        return {"err": f"시간 넘음({COM_TIMEOUT_S}초)"}
    line = next((l for l in r.stdout.splitlines() if l.startswith("RESULT ")), None)
    res = json.loads(line[7:]) if line else {"err": (r.stderr or r.stdout)[-400:]}
    res["sec"] = round(time.monotonic() - t0, 1)
    if dst.exists():
        import fitz
        d = fitz.open(str(dst))
        res["pdf_pages"] = d.page_count
        res["text"] = "".join(p.get_text() for p in d)
        d.close()
    return res


HWP_PER_FILE_S = 30        # 한글 일괄의 파일당 시간 — 보안 창을 아무도 안 누르면 COM_TIMEOUT_S + 이것 × 파일 수 뒤 끝낸다


def _hwp_pids() -> set:
    try:
        r = subprocess.run(["tasklist", "/FI", "IMAGENAME eq Hwp.exe", "/FO", "CSV", "/NH"], capture_output=True, text=True)
        return {int(x.split('","')[1]) for x in r.stdout.splitlines() if x.startswith('"Hwp.exe"')}
    except Exception:
        return set()


def com_hwp_batch(jobs, script: Path) -> list:
    """한글 하나로 여러 HWPX → PDF. `jobs=[(src, dst)]` → 결과 dict 목록(같은 차례)."""
    if not jobs:
        return []
    t0 = time.monotonic()
    args = [str(a) for s, d in jobs for a in (s, d)]
    before = _hwp_pids()
    try:
        r = subprocess.run([sys.executable, str(script), "hwp"] + args, capture_output=True, text=True,
                           encoding="utf-8", errors="replace", timeout=COM_TIMEOUT_S + HWP_PER_FILE_S * len(jobs))
        out = [json.loads(l[7:]) for l in r.stdout.splitlines() if l.startswith("RESULT ")]
    except subprocess.TimeoutExpired as e:
        out = [json.loads(l[7:]) for l in (e.stdout or b"").decode("utf-8", "replace").splitlines() if l.startswith("RESULT ")]
        for pid in _hwp_pids() - before:                 # 보안 창에서 멈춘 한글 — 이 시험이 띄운 것만 끝낸다
            subprocess.run(["taskkill", "/PID", str(pid), "/F"], capture_output=True)
    sec = round((time.monotonic() - t0) / max(1, len(jobs)), 1)
    res = []
    import fitz
    for k, (s, d) in enumerate(jobs):
        o = next((x for x in out if x.get("i") == k), {"err": "한글이 결과를 내지 않음(시간 넘음·멈춤)"})
        o["sec"] = sec
        if Path(d).exists():
            doc = fitz.open(str(d))
            o["pdf_pages"] = doc.page_count
            o["text"] = "".join(p.get_text() for p in doc)
            doc.close()
        res.append(o)
    return res


COM_OF = {"docx_edit": "word", "docx_shape": "word", "docx_word": "word", "hwpx_edit": "hwp", "hwpx_shape": "hwp", "pptx_shape": "ppt"}


# ── 압축·가림·비교 ──────────────────────────────────────────
def check_compress(src: Path) -> dict:
    import fitz
    from viewer import pdf_compress
    try:
        dst, a, b = pdf_compress.run(str(src), "mid", lambda *x: None)
    except pdf_compress.Skipped as e:
        return {"ok": True, "note": "건너뜀: " + str(e)}
    d0, d1 = fitz.open(str(src)), fitz.open(str(dst))
    try:
        same_pages = d0.page_count == d1.page_count
        n = min(5, d0.page_count)
        r = recall(words("".join(d0[i].get_text() for i in range(n))), "".join(d1[i].get_text() for i in range(n)))
    finally:
        d0.close(); d1.close()
    ok = same_pages and r >= 0.99 and b < a
    return {"ok": ok, "note": f"{a / 1e6:.1f}→{b / 1e6:.1f} MB ({100 * (a - b) / a:.0f}%), 쪽 같음 {same_pages}, 글 재현 {r:.3f}"}


def check_redact(src: Path) -> dict:
    import fitz
    from viewer import redact
    d = fitz.open(str(src))
    try:
        n = min(5, d.page_count)
        # 검색어는 원문 그대로의 낱말(정규화한 `words` 가 아니다 — O→0 처럼 바꾼 글자는 원본에서 찾을 수 없다)
        cand = sorted({w for w in re.findall(r"[A-Za-z가-힣]{4,}", "".join(d[i].get_text() for i in range(n)))},
                      key=len, reverse=True)
    finally:
        d.close()
    term = next((w for w in cand if len(w) >= 4 and w.isalpha()), None)
    if not term:
        return {"ok": True, "note": "앞 5쪽에 글자층 낱말이 없음 — 건너뜀"}
    marks = [m for m in redact.find_marks(str(src), keys=("phone", "email"), custom=term) if m["page"] < n]
    hits = [m for m in marks if m["kind"] == redact.KIND_SEARCH]
    if not hits:
        return {"ok": False, "note": f"검색어 '{term}' 을 찾지 못함"}
    dst, left = redact.apply(str(src), marks)
    d = fitz.open(str(dst))
    try:
        remain = sum(d[i].get_text().lower().count(term.lower()) for i in range(n))
    finally:
        d.close()
    return {"ok": remain == 0 and not left, "note": f"'{term}' {len(hits)}곳·패턴 {len(marks) - len(hits)}곳 가림, 결과에 남은 수 {remain}, 확인 남음 {len(left)}"}


def check_compare(src: Path, work: Path) -> dict:
    import fitz
    from viewer import compare
    d = fitz.open(str(src))
    n = d.page_count
    pg = min(1, n - 1)
    new = work / (src.stem + "_v2.pdf")
    d[pg].insert_text((72, 60), "PolyPDF sample check inserted line 7319", fontsize=11)
    d.save(str(new))
    d.close()
    t = time.time()
    os.utime(src, (t - 60, t - 60)); os.utime(new, (t, t))
    ch = compare.compare(str(src), str(new))
    ins = [c for c in ch if c["kind"] in ("insert", "replace") and any("7319" in x[2] for x in c["new"])]
    ok = bool(ins) and all(x[0] == pg for c in ins for x in c["new"])
    return {"ok": ok and len(ch) <= 3, "note": f"바뀐 곳 {len(ch)}곳, 더한 줄 찾음 {bool(ins)}(쪽 {pg + 1})"}


# ── 실행 ────────────────────────────────────────────────────
def main(argv) -> int:
    sys.stdout.reconfigure(encoding="utf-8")
    from PyQt6.QtCore import QStandardPaths
    QStandardPaths.setTestModeEnabled(True)
    from PyQt6.QtWidgets import QApplication
    app = QApplication.instance() or QApplication(sys.argv[:1])   # noqa: F841 — 변환이 Qt 이미지 도구를 쓴다
    import fitz
    from viewer import __version__, convert

    def opt(name, default=None):
        return argv[argv.index(name) + 1] if name in argv else default
    suite_p = Path(opt("--suite", str(HERE.parent / "sample_check_suite.json")))
    suite = json.loads(suite_p.read_text(encoding="utf-8"))
    only = set(filter(None, (opt("--only", "") or "").split(",")))
    keys = (opt("--formats") or "").split(",") if opt("--formats") else [f.key for f in convert.ready_formats()]
    use_com = "--no-com" not in argv
    stamp = time.strftime("%y%m%d-%H%M")
    out_dir = HERE.parent / "_review" / f"samples_{__version__}_{stamp}"
    work = out_dir / "work"
    work.mkdir(parents=True, exist_ok=True)
    script = out_dir / "_com.py"
    script.write_text(COM_SCRIPT, encoding="utf-8")
    have = {k: com_installed(k) for k in ("word", "hwp", "ppt")} if use_com else {}
    rows, fails = [], []

    def rec(case, what, ok, note, sec=None):
        rows.append({"case": case, "what": what, "ok": ok, "note": note, "sec": sec})
        print(("PASS" if ok else ("FAIL" if ok is False else "----")), case, what, note, flush=True)
        if ok is False:
            fails.append(f"{case} {what}: {note}")

    hwp_jobs = []

    def judge(job, res):
        case, key, prog, _out, _dst, sw_, has_text_, npg = job
        if res.get("err"):
            rec(case, key + " ▸ " + prog, False, res["err"], res.get("sec"))
            return
        rr = recall(sw_, res.get("text", "")) if has_text_ else None
        n_out = res.get("pdf_pages") or res.get("pages")
        note = f"{prog} 쪽 {n_out}" + (f" · 글 재현 {rr:.2f}" if rr is not None else "")
        ok = rr is None or rr >= RECALL_MIN[key] - 0.05
        if not ok and key in INFO_ONLY:
            ok, note = None, note + " (참고 — Word 자체 변환)"
        if key in ("hwpx_shape", "pptx_shape") and n_out != npg:
            ok, note = False, note + f" — 원본 {npg}쪽과 다름"
        rec(case, key + " ▸ " + prog, ok, note, res.get("sec"))

    if opt("--hwp-from"):
        # 한글 단계만 다시 — 앞선 실행의 HWPX(편집 가능·모양 그대로)를 한글 하나로 열어 판정한다(보안 창을 못 눌러 시간이 넘었을 때)
        base = Path(opt("--hwp-from"))
        (base / "_com.py").write_text(COM_SCRIPT, encoding="utf-8")
        jobs = []
        for c in suite["cases"]:
            cdir = base / "work" / c["name"]
            src = cdir / "sample.pdf"
            if not src.exists():
                continue
            d = fitz.open(str(src))
            pages = convert.parse_range(c.get("pages") or f"1-{min(10, d.page_count)}", d.page_count)
            txt = "".join(d[p].get_text() for p in pages)
            d.close()
            hx = sorted(cdir.glob("sample*.hwpx"), key=lambda f: f.stat().st_mtime)    # 만든 차례 = 편집 가능, 모양 그대로
            for f, key in zip(hx, ("hwpx_edit", "hwpx_shape")):
                jobs.append((c["name"], key, "hwp", f, f.with_suffix(".hwp.pdf"), words(txt), len(txt.strip()) > 50, len(pages)))
        print(f"== 한글로 {len(jobs)}개 열기 — 처음 뜨는 '접근 허용' 창에서 '모두 허용' 을 한 번 누르면 된다", flush=True)
        for job, res in zip(jobs, com_hwp_batch([(j[3], j[4]) for j in jobs], base / "_com.py")):
            judge(job, res)
        md = [f"# 표본 시험 — 한글 단계 ({stamp})", "", "| 표본 | 무엇 | 판정 | 내용 |", "| --- | --- | --- | --- |"]
        md += [f"| {r['case']} | {r['what']} | {'통과' if r['ok'] else ('**실패**' if r['ok'] is False else '—')} | {r['note']} |" for r in rows]
        md += ["", f"**판정: {'실패 ' + str(len(fails)) + '건' if fails else '통과'}**"]
        (base / "summary_hwp.md").write_text("\n".join(md), encoding="utf-8")
        shutil.rmtree(out_dir, ignore_errors=True)            # 이번에 만든 빈 실행 폴더
        print(md[-1], "\n결과:", base / "summary_hwp.md", flush=True)
        return 1 if fails else 0

    for c in suite["cases"]:
        if only and c["name"].split("_")[0] not in only and c["name"] not in only:
            continue
        src0 = Path(suite["samples"]) / c["file"]
        if not src0.exists():
            rec(c["name"], "표본", None, f"없음: {c['file']}")
            continue
        cdir = work / c["name"]
        cdir.mkdir(exist_ok=True)
        src = cdir / ("sample" + src0.suffix)
        shutil.copy2(src0, src)
        d = fitz.open(str(src))
        pages = convert.parse_range(c.get("pages") or f"1-{min(10, d.page_count)}", d.page_count)
        src_text = "".join(d[p].get_text() for p in pages)
        has_text = len(src_text.strip()) > 50
        n_pages = d.page_count
        d.close()
        sw = words(src_text)
        print(f"\n== {c['name']} ({n_pages}쪽, 시험 {len(pages)}쪽, 글자층 {'있음' if has_text else '없음'})", flush=True)
        for key in keys:
            t0 = time.monotonic()
            try:
                out = convert.run(str(src), key, pages)
            except convert.NothingToWrite:
                rec(c["name"], key, True if key == "xlsx" else False, "만들 것 없음(표 없음)", round(time.monotonic() - t0, 1))
                continue
            except Exception as e:                       # noqa: BLE001
                if key in INFO_ONLY and "그림으로만" in str(e):  # Word 가 쪽을 그림으로만 바꿈 — 앱이 알리는 정상 동작(형식 변환 SOT §4.4)
                    rec(c["name"], key, None, "Word 가 그림으로만 바꿈 — 앱이 다른 방식을 알림(참고)", round(time.monotonic() - t0, 1))
                    continue
                rec(c["name"], key, False, f"변환 실패: {e!r}"[:300] + "\n" + traceback.format_exc()[-600:])
                continue
            sec = round(time.monotonic() - t0, 1)
            try:
                info = lib_read(key, Path(out))
            except Exception as e:                       # noqa: BLE001
                rec(c["name"], key, False, f"라이브러리로 열기 실패: {e!r}"[:300], sec)
                continue
            r = recall(sw, info.get("text", "")) if has_text and key in RECALL_MIN else None
            shape = {k: v for k, v in info.items() if k != "text"}
            ok = True
            note = f"{shape}"
            if r is not None:
                note += f" 글 재현 {r:.2f}"
                ok = r >= RECALL_MIN[key]
                if not ok and key in INFO_ONLY:
                    ok, note = None, note + " (참고 — Word 자체 변환)"
            if key == "pptx_shape" and info.get("slides") != len(pages):
                ok, note = False, note + f" — 슬라이드 수 ≠ 쪽 수({len(pages)})"
            if key == "png" and info.get("files") != len(pages):
                ok, note = False, note + f" — 그림 수 ≠ 쪽 수({len(pages)})"
            rec(c["name"], key, ok, note, sec)
            prog = COM_OF.get(key)
            if prog and use_com and have.get(prog):
                job = (c["name"], key, prog, Path(out), Path(out).with_suffix(f".{prog}.pdf"), sw, has_text, len(pages))
                if prog == "hwp":
                    hwp_jobs.append(job)                 # 한글은 끝에 한 번에(보안 창을 한 번만)
                else:
                    judge(job, com_to_pdf(prog, job[3], job[4], script))
        if n_pages <= SMALL_PAGES:
            for what, fn in (("압축", lambda: check_compress(src)), ("가림", lambda: check_redact(src)),
                             ("비교", lambda: check_compare(src, cdir))):
                t0 = time.monotonic()
                try:
                    r = fn()
                    rec(c["name"], what, r["ok"], r["note"], round(time.monotonic() - t0, 1))
                except Exception as e:                   # noqa: BLE001
                    rec(c["name"], what, False, f"{e!r}"[:300])

    if hwp_jobs:
        print(f"\n== 한글로 {len(hwp_jobs)}개 열기 — 처음 뜨는 '접근 허용' 창에서 '모두 허용' 을 한 번 누르면 된다", flush=True)
        for job, res in zip(hwp_jobs, com_hwp_batch([(j[3], j[4]) for j in hwp_jobs], script)):
            judge(job, res)

    md = [f"# 표본 시험 — {__version__} ({stamp})", "",
          f"- 표본 `{suite_p.name}` · 형식 {', '.join(keys)} · COM {'켬 ' + str(have) if use_com else '끔'}", "",
          "| 표본 | 무엇 | 판정 | 초 | 내용 |", "| --- | --- | --- | --- | --- |"]
    for r in rows:
        mark = "통과" if r["ok"] else ("**실패**" if r["ok"] is False else "—")
        md.append(f"| {r['case']} | {r['what']} | {mark} | {r['sec'] if r['sec'] is not None else ''} | "
                  f"{str(r['note']).splitlines()[0].replace('|', '/')} |")
    md += ["", f"**판정: {'실패 ' + str(len(fails)) + '건' if fails else '통과'}**"]
    (out_dir / "summary.md").write_text("\n".join(md), encoding="utf-8")
    (out_dir / "results.json").write_text(json.dumps(rows, ensure_ascii=False, indent=1), encoding="utf-8")
    print("\n".join(md[-1:]), "\n결과:", out_dir, flush=True)
    return 1 if fails else 0


if __name__ == "__main__":
    code = main(sys.argv[1:])
    sys.stdout.flush()
    os._exit(code)
