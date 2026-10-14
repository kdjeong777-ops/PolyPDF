# -*- coding: utf-8 -*-
"""261014-30: 표본 시험(릴리스 SOT §4.4 — 실제 업무 PDF)에서 나온 변환·가림 결함을 합성 PDF 로 고정한다.

A. 테두리 친 문단(선 넷 + 긴 글 줄)은 그림이 아니라 글이다 — 글 줄이 남는다 (형식 변환 SOT §1, `FIG_TEXT_LINES`)
B. 2단 — 단마다 줄 높이가 몇 pt 어긋나도 나란한 줄로 묶는다(따로 쌓여 쪽이 넘쳤다, `BAND_OVERLAP`)
C. 글이 든 열이 하나뿐인 '표' 는 표가 아니다(2단의 단 하나가 1열 표로 잡혀 본문이 표 칸에 들어갔다, 텍스트 창 SOT §3 규칙 9)
D. 머리말 홀·짝 — 문서 여럿을 이은 PDF 는 홀·짝으로 나누지 않는다, 진짜 홀·짝은 나눈다(`EDGE_PARITY_MIN`)
E. 표 칸 글은 원본 크기로(HWPX 모양 그대로·PowerPoint — 큰 표가 원본보다 길어져 쪽을 넘었다)
F. 가림 — 옆 줄과 글자 상자가 겹쳐도 그 줄 글을 지우지 않고, 목록·'남은 글' 확인에도 넣지 않는다(보안 SOT §14)
G. Word(Microsoft Word 사용)가 쪽을 그림으로만 바꾸면 파일을 지우고 다른 방식을 알린다(형식 변환 SOT §4.4)
H. HWPX 모양 그대로 — 꼬리말 자리에 남은 본문 줄(꼬리말 찾기가 쪽 번호만 잡은 줄)은 꼬리말로 옮긴다(넘쳐 쪽이 늘었다)
I. 글자층의 제어 문자(글꼴이 대시를 제어 문자로 대응)가 있어도 Word·한글·PowerPoint 쓰기가 실패하지 않는다(`source.xml_safe`)
J. 줄 상자보다 큰 글자 크기(OCR 잡음)는 상자 높이로(줄 높이가 아래 줄을 밀어 쪽이 늘었다)
"""
import os, sys, tempfile, shutil, zipfile, re
os.environ["QT_QPA_PLATFORM"] = "offscreen"
sys.stdout.reconfigure(encoding="utf-8")
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from pathlib import Path

import fitz
from PyQt6.QtCore import QStandardPaths
QStandardPaths.setTestModeEnabled(True)
from PyQt6.QtWidgets import QApplication

app = QApplication.instance() or QApplication(sys.argv[:1])
fails = []


def chk(cond, msg, extra=""):
    print(("PASS" if cond else "FAIL"), "-", msg, extra if not cond else "", flush=True)
    if not cond:
        fails.append(msg)


root = Path(tempfile.mkdtemp(prefix="polypdf_convfix_test_"))
try:
    from viewer.convert.source import PageSource
    from viewer.convert import layout
    from viewer import convert

    # ── A. 테두리 친 문단 ─────────────────────────────────────────
    a = root / "A.pdf"
    d = fitz.open(); p = d.new_page()
    # 실제 지침 PDF(표본 S1 3쪽)처럼 — 테두리 선 넷이 따로따로 + 회색 제목 띠
    for seg in (((60, 100), (500, 100)), ((500, 100), (500, 220)), ((500, 220), (60, 220)), ((60, 220), (60, 100))):
        sh = p.new_shape(); sh.draw_line(*seg); sh.finish(color=(0, 0, 0), width=1); sh.commit()
    sh = p.new_shape(); sh.draw_rect(fitz.Rect(150, 104, 400, 118)); sh.finish(fill=(0.84, 0.84, 0.84), color=None); sh.commit()
    body = ["This guide applies from the publication date to design services and",
            "construction works already under way when the owner finds it necessary.",
            "However the quality standard for recycled cold mixtures is deferred."]
    for k, t in enumerate(body):
        p.insert_text((70, 130 + 25 * k), t, fontsize=11)
    d.save(str(a)); d.close()
    with PageSource(str(a), tables="off") as ps:
        bl = layout.page_blocks(ps, 0, join=False)
    got = " ".join(b["text"] for b in bl if b["kind"] == "line")
    chk(all(t.split()[0] in got for t in body) and not any(b["kind"] == "image" for b in bl),
        "A 테두리 친 문단은 글로 남는다(그림으로 삼키지 않는다)", str([b["kind"] for b in bl]))

    # ── B. 2단, 단마다 줄 높이 어긋남 ────────────────────────────────
    b2 = root / "B.pdf"
    d = fitz.open(); p = d.new_page()
    for k in range(20):
        p.insert_text((50, 100 + 14 * k), f"left column line number {k:02d} text", fontsize=10)
        p.insert_text((310, 105 + 14 * k), f"right column line number {k:02d} text", fontsize=10)
    d.save(str(b2)); d.close()
    with PageSource(str(b2), tables="off") as ps:
        lines = [b for b in layout.page_blocks(ps, 0, join=False) if b["kind"] == "line"]
    bs = layout.bands(lines)
    chk(len(lines) == 40 and len(bs) == 20, "B 2단 — 줄 높이가 5 pt 어긋나도 나란한 줄로 묶는다(쌓지 않는다)",
        f"lines={len(lines)} bands={len(bs)}")

    # ── C. 1열 '표' ──────────────────────────────────────────────
    from viewer import text_extract2 as tx
    one = [((0, 0, 100, 100), [["first paragraph line"], ["second line"], [None]])]
    two = [((0, 0, 100, 100), [["a", "b"], ["c", None]])]
    chk(tx._real_tables(one) == [] and tx._real_tables(two) == two, "C 글이 든 열이 하나면 표가 아니다(둘이면 표)")
    c1 = root / "C.pdf"
    d = fitz.open(); p = d.new_page()
    sh = p.new_shape()
    for k in range(6):                                           # 단 하나를 가로줄로 나눈 상자(1열 '표')
        sh.draw_line((300, 100 + 20 * k), (540, 100 + 20 * k))
    sh.draw_line((300, 100), (300, 200)); sh.draw_line((540, 100), (540, 200))
    sh.finish(color=(0, 0, 0), width=0.5); sh.commit()
    for k in range(5):
        p.insert_text((305, 114 + 20 * k), f"column body sentence {k}", fontsize=10)
    d.save(str(c1)); d.close()
    with PageSource(str(c1), tables="off") as ps:
        kinds = [b["kind"] for b in layout.page_blocks(ps, 0, join=False)]
    chk("table" not in kinds and kinds.count("line") == 5, "C 1열 '표' 의 글은 줄로 남는다", str(kinds))

    # ── D. 머리말 홀·짝 ───────────────────────────────────────────
    def edge(text, where="head"):
        return {"where": where, "text": text, "pieces": [(72, 300, text)], "rect": (72, 30, 300, 40)}
    merged = {i: [edge("BMD-001 NAPA" if i < 4 else ("BMD-002 NAPA" if i < 8 else "BMD-004 NAPA"))] for i in range(10)}
    pm = layout.edge_plan(merged)
    chk(list(pm["head"]) == ["BOTH"], "D 문서 여럿을 이은 PDF 는 홀·짝으로 나누지 않는다", str(list(pm["head"])))
    oddeven = {i: [edge("Chapter Title" if i % 2 else "Book Name")] for i in range(10)}
    po = layout.edge_plan(oddeven)
    chk(set(po["head"]) == {"EVEN", "ODD"}, "D 진짜 홀·짝은 나눈다", str(list(po["head"])))

    # ── E. 표 칸 글 크기 ─────────────────────────────────────────
    e = root / "E.pdf"
    d = fitz.open(); p = d.new_page()
    sh = p.new_shape()
    for k in range(5):
        sh.draw_line((72, 100 + 18 * k), (472, 100 + 18 * k))
    for x in (72, 222, 372, 472):
        sh.draw_line((x, 100), (x, 172))
    sh.finish(color=(0, 0, 0), width=0.5); sh.commit()
    for k in range(4):
        for j, x in enumerate((76, 226, 376)):
            p.insert_text((x, 113 + 18 * k), f"cell {k}{j} value", fontsize=7)
    d.save(str(e)); d.close()
    with PageSource(str(e), tables="off") as ps:
        tb = [b for b in layout.page_blocks(ps, 0, join=False) if b["kind"] == "table"]
    chk(tb and abs(tb[0]["size"] - 7.0) < 0.6, "E 표 블록에 원본 글 크기", str([t.get("size") for t in tb]))
    out = convert.run(str(e), "hwpx_shape")
    with zipfile.ZipFile(out) as z:
        hdr = z.read("Contents/header.xml").decode("utf-8")
        sec = z.read("Contents/section0.xml").decode("utf-8")
    tbl = sec[sec.find("<hp:tbl"):sec.rfind("</hp:tbl>")]
    refs = set(re.findall(r'<hp:run charPrIDRef="(\d+)"', tbl))
    heights = {}
    for m in re.finditer(r'<hh:charPr\s[^>]*>', hdr):
        i, h = re.search(r'\sid="(\d+)"', m.group(0)), re.search(r'\sheight="(\d+)"', m.group(0))
        if i and h:
            heights[i.group(1)] = int(h.group(1))
    chk(refs and all(heights.get(r) == 700 for r in refs), "E HWPX 모양 그대로 — 표 칸 글이 원본 크기(7 pt)",
        f"{refs} {[heights.get(r) for r in refs]}")
    import pptx
    outp = convert.run(str(e), "pptx_shape")
    sizes = {r.font.size.pt for s in pptx.Presentation(str(outp)).slides for shp in s.shapes if shp.has_table
             for row in shp.table.rows for c in row.cells for para in c.text_frame.paragraphs for r in para.runs}
    chk(sizes and max(sizes) < 9, "E PowerPoint — 표 칸 글이 원본 크기에 맞춘다(10 pt 고정이 아니다)", str(sizes))

    # ── F. 가림 — 1 pt 겹친 윗줄 ───────────────────────────────────
    f = root / "F.pdf"
    d = fitz.open(); p = d.new_page()
    p.insert_text((72, 100), "Gediminas Technical University,", fontsize=11)
    p.insert_text((72, 110.5), "hong.gd@example.com", fontsize=11)              # 윗줄과 세로로 살짝 겹친다
    d.save(str(f)); d.close()
    from viewer import redact
    marks = redact.find_marks(str(f), keys=("email",))
    dst, left = redact.apply(str(f), marks)
    chk([m["text"] for m in marks] == ["hong.gd@example.com"] and left == [] and
        "Technical" in fitz.open(str(dst))[0].get_text(),
        "F 가림 목록·확인은 그 자리 글만(윗줄은 넣지 않고 지우지도 않는다)", f"{[m['text'] for m in marks]} {left}")

    # ── G. Word 가 쪽을 그림으로만 바꾸면 '변환했습니다' 로 내지 않는다 ─────────────
    from viewer.convert import word_com
    import docx as _docx
    _real = word_com.convert_word

    def _images_only(ps, pages, dst, tick):
        dd = _docx.Document(); dd.add_paragraph(""); dd.save(str(dst))      # 글 없는 결과(그림만 — 표본 S9 의 Word 결과)
    word_com.convert_word = _images_only
    try:
        err = None
        try:
            convert.run(str(b2), "docx_word")
        except RuntimeError as ex:
            err = str(ex)
        chk(err and "그림으로만" in err and not (root / "B.docx").exists(),
            "G Word 결과에 글이 없으면 지우고 다른 방식을 알린다", str(err))
    finally:
        word_com.convert_word = _real

    # ── H. 꼬리말 자리 안에 남은 본문 줄은 꼬리말로(HWPX 모양 그대로 — 본문 영역 밖이라 새 쪽으로 넘쳤다) ─────
    from viewer.convert import hwpx_writer as hw
    foot = [{"where": "foot", "rect": (460, 681.6, 462.4, 692.1), "pieces": [(460, 462.4, "i")], "text": "i",
             "size": 8, "bold": False, "serif": True, "pw": 595, "ph": 737}]
    body = [{"kind": "line", "rect": (68, 600, 470, 611), "text": "본문 줄", "size": 10.6, "bold": False, "serif": True},
            {"kind": "line", "rect": (313.1, 683.5, 437.8, 691.5), "text": "아스팔트 콘크리트 포장 시공 지침", "size": 8,
             "bold": False, "serif": True}]
    kept = hw._edge_zone_lines(body, [], foot, 595, 737)
    chk([k["text"] for k in kept] == ["본문 줄"] and [p_[2] for p_ in foot[0]["pieces"]] == ["아스팔트 콘크리트 포장 시공 지침", "i"],
        "H 꼬리말 자리의 본문 줄은 같은 꼬리말 줄의 조각으로 옮긴다", f"{[k['text'] for k in kept]} {foot[0]['pieces']}")

    # ── I. 글자층의 제어 문자(글꼴이 대시를  로 대응) — Word·한글·PowerPoint 쓰기가 실패하지 않는다 ─────────
    ic = root / "I.pdf"
    d = fitz.open(); p = d.new_page()
    p.insert_text((72, 100), "temperature range 25" + chr(1) + "30 C under traffic loads", fontsize=11)
    d.save(str(ic)); d.close()
    okk = []
    for k in ("docx_edit", "hwpx_edit", "pptx_shape", "docx_shape"):
        try:
            convert.run(str(ic), k)
            okk.append(k)
        except Exception as ex:
            print("   ", k, repr(ex)[:120])
    import docx as _dx
    dt = " ".join(x.text for x in _dx.Document(str(root / "I.docx")).paragraphs) if (root / "I.docx").exists() else ""
    chk(len(okk) == 4 and "25" + chr(0xFFFD) + "30" in dt, "I 제어 문자가 든 글자층도 변환된다(알 수 없는 글자 표시로)", f"{okk} {dt!r}")

    # ── J. 줄 상자보다 큰 글자 크기(스캔 표지의 OCR 잡음)는 상자 높이로 — 줄 높이가 아래 줄을 밀어 쪽이 늘었다 ─────
    band = [{"rect": (179, 328, 320, 501), "size": 221.1, "text": "X", "bold": False, "serif": True}]
    normal = [{"rect": (72, 100, 300, 114), "size": 11.0, "text": "body", "bold": False, "serif": True}]
    chk(layout.band_text(band)[3] == 173 and layout.band_text(normal)[3] == 11.0,
        "J 글자 크기는 줄 상자 높이를 넘지 않는다(보통 줄은 그대로)",
        f"{layout.band_text(band)[3]} {layout.band_text(normal)[3]}")
finally:
    shutil.rmtree(root, ignore_errors=True)

print("\n=== " + ("ALL PASS" if not fails else f"FAILURE ({len(fails)})") + " ===")
for m in fails:
    print(" -", m)
sys.stdout.flush()
os._exit(1 if fails else 0)
