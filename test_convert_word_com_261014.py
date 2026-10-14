# -*- coding: utf-8 -*-
"""261014-20: Word(Microsoft Word 사용) 변환 (형식 변환 SOT §4.4, 개선 로드맵 P3-2).

Word 가 없는 PC 에서는 A·F 만 본다(나머지는 SKIP — Word 가 있는 PC 에서 의미가 있다).
A. Word 가 있으면 형식 목록에 있고, F. 없으면 없다
B. 실제 처리기로 변환 — pdfreflow 안내 창에 막히지 않는다, docx 에 글
C. Word 설정(`DisableConvertPdfWarning`)을 바꾸지 않는다
D. 고른 쪽만
E. 중지하면 그 Word 를 끝내고 만들던 docx 를 지운다
G. 끝나면 이 변환이 띄운 Word·pdfreflow 가 남지 않는다(사용자 Word 는 건드리지 않는다)
"""
import os, sys, tempfile, shutil, time
os.environ["QT_QPA_PLATFORM"] = "offscreen"
sys.stdout.reconfigure(encoding="utf-8")
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from pathlib import Path

import fitz
from PyQt6.QtCore import QStandardPaths
QStandardPaths.setTestModeEnabled(True)
from PyQt6.QtWidgets import QApplication, QMessageBox
from PyQt6.QtTest import QTest

app = QApplication.instance() or QApplication(sys.argv[:1])
fails = []


def chk(cond, msg, extra=""):
    print(("PASS" if cond else "FAIL"), "-", msg, extra if not cond else "", flush=True)
    if not cond:
        fails.append(msg)


from viewer import convert
from viewer.convert import word_com

word_com._CACHE["installed"] = False
chk("docx_word" not in [f.key for f in convert.ready_formats()], "F Word 가 없으면 형식 목록에 없다")
word_com._CACHE.pop("installed", None)
have = word_com.word_installed()
chk(("docx_word" in [f.key for f in convert.ready_formats()]) == have, "A Word 가 있으면 형식 목록에 있다", str(have))


def opt():
    try:
        import winreg
        with winreg.OpenKey(winreg.HKEY_CURRENT_USER, r"Software\Microsoft\Office\16.0\Word\Options") as k:
            return winreg.QueryValueEx(k, "DisableConvertPdfWarning")[0]
    except Exception:
        return None


def mine():
    return word_com._pids("winword.exe") | word_com._pids("pdfreflow.exe")


root = Path(tempfile.mkdtemp(prefix="polypdf_wordcom_test_"))
try:
    if not have:
        print("SKIP - B~G (Word 가 없는 PC)")
    else:
        src = root / "W.pdf"
        d = fitz.open()
        for i in range(3):
            d.new_page().insert_text((72, 100), f"Word page {i + 1} sentence for conversion.", fontsize=12)
        d.save(str(src)); d.close()
        opt0, pids0 = opt(), mine()

        QMessageBox.question = staticmethod(lambda *x, **k: QMessageBox.StandardButton.No)
        infos = []
        QMessageBox.information = staticmethod(lambda *x, **k: infos.append(x[2] if len(x) > 2 else ""))
        QMessageBox.warning = staticmethod(lambda *x, **k: infos.append(x[2] if len(x) > 2 else ""))
        from viewer.widgets import convert_dialog as cdmod
        cdmod.ConvertDialog.exec = lambda dlg: (dlg._accept(), cdmod.ConvertDialog.DialogCode.Accepted)[1]
        from viewer.app import MainWindow
        from viewer.history import HistoryItem
        mw = MainWindow(); mw.show(); app.processEvents()
        mw.open_folder(root); QTest.qWait(300)
        mw._load_main(HistoryItem(str(src), 0, "", "bookmark")); QTest.qWait(300)

        t0 = time.perf_counter()
        ok = mw.action_convert(fmt="docx_word"); QTest.qWait(300)
        took = time.perf_counter() - t0
        out = root / "W.docx"
        from docx import Document
        txt = "\n".join(p.text for p in Document(str(out)).paragraphs) if out.exists() else ""
        chk(ok and "Word page 2 sentence" in txt and took < 120, "B 실제 처리기로 Word 변환(안내 창에 막히지 않는다)",
            f"{took:.1f}s {txt[:80]} {infos[-1:]}")
        chk(opt() == opt0, "C Word 설정을 바꾸지 않는다", f"{opt0} → {opt()}")
        time.sleep(1.0)
        chk(not (mine() - pids0), "G 이 변환이 띄운 Word·pdfreflow 가 남지 않는다", str(mine() - pids0))

        out2 = convert.run(str(src), "docx_word", pages=[1])
        t2 = "\n".join(p.text for p in Document(str(out2)).paragraphs)
        chk("Word page 2" in t2 and "Word page 1" not in t2 and "Word page 3" not in t2, "D 고른 쪽만", t2[:120])

        n = {"k": 0}

        def stop_soon(i, total, label):
            n["k"] += 1
            return n["k"] < 2
        try:
            convert.run(str(src), "docx_word", progress=stop_soon)
            stopped = False
        except convert.Cancelled:
            stopped = True
        time.sleep(1.0)
        made = sorted(p.name for p in root.glob("*.docx"))
        chk(stopped and not (mine() - pids0) and len(made) == 2,
            "E 중지하면 그 Word 를 끝내고 만들던 docx 를 지운다", f"{stopped} {mine() - pids0} {made}")
finally:
    shutil.rmtree(root, ignore_errors=True)

print("\n=== " + ("ALL PASS" if not fails else f"FAILURE ({len(fails)})") + " ===")
for m in fails:
    print(" -", m)
sys.stdout.flush()
os._exit(1 if fails else 0)
