# -*- coding: utf-8 -*-
"""260908-2: 텍스트 창 내용을 **Word 로 저장** (텍스트 창 SOT §7).

- 범위는 부르는 쪽이 정한다 — 현재 쪽 / 쪽 범위 / 문서 전체(사용자 결정 260908).
- 문단 스타일은 창에서 정한 값(글꼴·크기·굵기·자간·줄간격)을 그대로 얹는다.
  제목은 `Heading 1`, 내용은 `Normal` 위에 올린다 — Word 의 목차·탐색창이 그대로 먹는다.
- 표는 **글줄**로 나간다(SOT §3.3·§7). 표 개체로 만들면 원본 구조를 반쯤 맞게 복원해
  오히려 고치기 번거롭다.
"""
from __future__ import annotations


def _apply_run(run, st):
    from docx.shared import Pt
    run.font.name = st.get("family", "맑은 고딕")
    run.font.size = Pt(float(st.get("size_pt", 11)))
    run.bold = bool(st.get("bold", False))
    try:                                    # 자간(pt) — 한글 글꼴 포함 동작
        from docx.oxml.ns import qn
        from docx.oxml import OxmlElement
        sp = float(st.get("spacing_pt", 0.0) or 0.0)
        if sp:
            el = OxmlElement("w:spacing")
            el.set(qn("w:val"), str(int(round(sp * 20))))   # twip(1/20 pt)
            run._element.get_or_add_rPr().append(el)
        # 한글 글꼴이 서양 글꼴로 밀리지 않게 eastAsia 도 지정
        rf = run._element.get_or_add_rPr().get_or_add_rFonts()
        rf.set(qn("w:eastAsia"), st.get("family", "맑은 고딕"))
    except Exception:
        pass


def export_pages_to_docx(pdf_path, doc_fitz, pages, dst, styles, *,
                         omit_tables: bool = False,
                         fix_rows=None, ocr_lookup=None, words_lookup=None,
                         progress=None) -> tuple:
    """쪽 목록을 docx 로. 반환 (성공여부, 메시지).

    `progress(했음, 전체, 글) -> bool` 이 False 를 주면 멈추고 `(False, "")` 를 돌려준다(취소 — 파일을 쓰지 않는다)."""
    try:
        from docx import Document
        from docx.shared import Pt
    except Exception as e:                       # noqa: BLE001
        return False, f"python-docx 없음: {e}"
    from viewer import text_extract2 as tx

    out = Document()
    try:
        for i, pg in enumerate(pages):
            if progress is not None and progress(i, len(pages), f"p.{pg + 1}") is False:
                return False, ""
            ocr = ""
            if not tx.has_text_layer(doc_fitz, pg):
                ocr = (ocr_lookup(pg) if ocr_lookup else "") or ""
            # 260913-6(텍스트 창 SOT §3.1·§7): 원천은 창과 같다 — 스캔 쪽은 우리 OCR
            ws, wdpi = ((words_lookup(pg) or ([], 0)) if words_lookup else ([], 0))
            rows = tx.page_lines(doc_fitz, pdf_path, pg,
                                 tables=("omit" if omit_tables else "lines"),
                                 ocr_text=ocr, ocr_words=ws or None, ocr_dpi=int(wdpi or 0))
            # 260913-4(텍스트 창 SOT §5.1.3): 줄 번호가 아니라 **자리로** 얹는다(합친 줄 포함).
            #   `fix_rows(쪽, 줄목록) -> 줄목록` — 창·본문 읽기·단어장과 같은 함수를 쓴다.
            if fix_rows is not None:
                try:
                    rows = fix_rows(pg, rows)
                except Exception:
                    pass
            if not rows:
                continue
            for r in rows:
                key = r.get("style", "body")
                st = (styles or {}).get(key, {})
                p = out.add_paragraph()
                if key == "title":
                    try:
                        p.style = out.styles["Heading 1"]
                    except Exception:
                        pass
                run = p.add_run(r.get("text", ""))
                _apply_run(run, st)
                try:
                    p.paragraph_format.line_spacing = float(
                        st.get("line_spacing", 1.2))
                except Exception:
                    pass
            if pg != pages[-1]:
                out.add_page_break()
        out.save(str(dst))
        return True, ""
    except Exception as e:                       # noqa: BLE001
        return False, str(e)


def export_docx_job(pdf_path, pages, dst, styles, *, omit_tables: bool = False,
                    fixes=None, forced_pages=(), progress=None) -> tuple:
    """261014-5(텍스트 창 SOT §7, 응답성 SOT §4): Word 저장을 **배경 스레드에서** — 문서·study.db 연결을 이 스레드에서 따로 연다.

    종전에는 메인 스레드가 바쁨 커서만 띄우고 돌아, '문서 전체' 는 쪽마다 표 인식까지 하느라 수 초~수 분 창이 멈췄다.
    PyMuPDF 문서와 sqlite 연결은 만든 스레드에서만 쓴다(응답성 SOT §4 ⑤) — 메인의 본문 문서·`_study_read` 연결을 쓰지 않는다.
    `fixes` 는 고침 저장소의 **사본**(메인이 떠서 준다 — 내보내는 동안 사용자가 고쳐도 섞이지 않게).
    `forced_pages` 는 [OCR 다시 읽기] 한 쪽(그 쪽은 글자층이 있어도 우리 OCR 을 쓴다, SOT §3.1)."""
    import fitz
    from viewer import text_extract2 as tx
    doc = fitz.open(str(pdf_path))
    st = None
    try:
        if doc.needs_pass:
            from viewer import secure_store
            doc.authenticate(secure_store.recall_any(str(pdf_path)) or "")
        try:
            from viewer.study.study_store import StudyStore, file_key_for
            from viewer import dbutil
            st = StudyStore(busy_ms=dbutil.BUSY_MS_UI)
            key = file_key_for(str(pdf_path))
        except Exception:
            st = None
            key = None

        def _text(pg):
            try:
                return (st.get_page_text(key, int(pg)) or "") if st else ""
            except Exception:
                return ""

        def _words(pg):
            if st is None:
                return ([], 0)
            forced = int(pg) in set(forced_pages or ())
            return tx.pick_ocr_words(
                doc, int(pg),
                lambda k: (st.get_page_words(key, int(k)), int(st.get_page_dpi(key, int(k)) or 0)),
                forced=forced)

        fix = (lambda pg, rows: fixes.apply_to_rows(str(pdf_path), pg, rows)) if fixes is not None else None
        return export_pages_to_docx(str(pdf_path), doc, list(pages), dst, styles, omit_tables=omit_tables,
                                    fix_rows=fix, ocr_lookup=_text, words_lookup=_words, progress=progress)
    finally:
        if st is not None:
            try:
                st.close()
            except Exception:
                pass
        doc.close()


def snapshot_fixes():
    """고침 저장소의 사본(배경 내보내기용) — 메인 스레드에서 부른다."""
    import copy
    from viewer.text_fix_store import TextFixStore, store
    src = store()
    snap = TextFixStore.__new__(TextFixStore)
    snap.path = src.path
    snap._data = copy.deepcopy(src._data)
    return snap
