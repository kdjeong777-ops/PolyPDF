# -*- coding: utf-8 -*-
"""PowerPoint 모양 그대로 — 고칠 수 있는 글상자 + 그림 + 표 (형식 변환 SOT §1, 261014-19) — python-pptx(MIT, §4.3).

쪽마다 슬라이드 하나. 슬라이드 크기는 문서에 하나뿐이라 첫 쪽 크기로 하고, 크기가 다른 쪽은 비율을 지켜 줄여 가운데에 놓는다.
글 줄은 HWPX 모양 그대로와 같은 묶음(`layout.bands`)마다 글상자 하나 — 원본 자리·크기·굵기·글꼴·글자 간격.
"""
from __future__ import annotations

import io
from pathlib import Path

from viewer.convert import layout

EMU_PER_PT = 12700
SPC_MIN, SPC_MAX = -400, 1000       # 글자 간격(1/100 pt) 범위 — 너무 좁히거나 벌리지 않게
TABLE_FONT_PT = 10                  # 표 칸 글 크기
BOX_SLACK_PT = 2                    # 글상자 너비 여유 — 재는 값과 그리는 값의 작은 차로 줄이 넘어가지 않게


def _emu(pt: float) -> int:
    return int(round(pt * EMU_PER_PT))


def _ea_font(run, face: str) -> None:
    """동아시아 글꼴(`a:ea`)도 같은 글꼴로 — 한글이 다른 글꼴로 바뀌지 않게."""
    from pptx.oxml.ns import qn
    rpr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea"):
        el = rpr.find(qn(tag))
        if el is None:
            el = rpr.makeelement(qn(tag), {})
            rpr.append(el)
        el.set("typeface", face)


def _spc(text: str, width: float, size: float, serif: bool) -> int:
    m = layout.measure(text, size, serif)
    if m is None or len(text) < 2 or width <= 0:
        return 0
    v = (width - m) / len(text) * 100.0
    return int(max(SPC_MIN, min(SPC_MAX, round(v))))


def _text_box(slide, place, text, x0, y0, width, height, size, bold, serif) -> None:
    from pptx.util import Emu, Pt
    left, top, scale = place
    tb = slide.shapes.add_textbox(Emu(_emu(left + x0 * scale)), Emu(_emu(top + y0 * scale)),
                                  Emu(_emu((width + BOX_SLACK_PT) * scale)), Emu(_emu(max(height, size) * scale)))
    tf = tb.text_frame
    tf.word_wrap = False
    tf.auto_size = None
    for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, side, 0)
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = text
    r.font.size = Pt(round(size * scale * 2) / 2)
    r.font.bold = bool(bold)
    face = layout.SERIF_FACE if serif else layout.SANS_FACE
    r.font.name = face
    _ea_font(r, face)
    spc = _spc(text, width, size, serif)
    if spc:
        r._r.get_or_add_rPr().set("spc", str(int(spc * scale)))


def _table(slide, place, b) -> None:
    from pptx.util import Emu, Pt
    from pptx.dml.color import RGBColor
    left, top, scale = place
    rows = b["rows"]
    x0, y0, x1, y1 = b["rect"]
    nr, nc = len(rows), max(len(r) for r in rows)
    gf = slide.shapes.add_table(nr, nc, Emu(_emu(left + x0 * scale)), Emu(_emu(top + y0 * scale)),
                                Emu(_emu((x1 - x0) * scale)), Emu(_emu((y1 - y0) * scale)))
    tbl = gf.table
    tbl.first_row = False
    tbl.horz_banding = False
    for ri in range(nr):
        tbl.rows[ri].height = Emu(_emu((y1 - y0) * scale / nr))
        for ci in range(nc):
            cell = tbl.cell(ri, ci)
            val = rows[ri][ci] if ci < len(rows[ri]) else None
            cell.text = "" if val is None else str(val)
            cell.fill.background()                       # 채움 없음 — 원본처럼
            cell.margin_left = cell.margin_right = Emu(_emu(2))
            cell.margin_top = cell.margin_bottom = Emu(_emu(1))
            for para in cell.text_frame.paragraphs:
                for r in para.runs:
                    r.font.size = Pt(TABLE_FONT_PT * scale)
                    r.font.color.rgb = RGBColor(0, 0, 0)
            _cell_border(cell)


def _cell_border(cell) -> None:
    """가는 검은 테두리(0.5 pt) — 표 서식이 칠한 파란 채움·흰 선 대신."""
    from pptx.oxml.ns import qn
    tcPr = cell._tc.get_or_add_tcPr()
    # 규격 차례(lnL·lnR·lnT·lnB 가 채움보다 앞) — 거꾸로 넣으면 파워포인트가 세로 선을 버린다(261014-19 대조)
    for k, tag in enumerate(("a:lnL", "a:lnR", "a:lnT", "a:lnB")):
        old = tcPr.find(qn(tag))
        if old is not None:
            tcPr.remove(old)
        ln = tcPr.makeelement(qn(tag), {"w": str(6350), "cap": "flat", "cmpd": "sng", "algn": "ctr"})
        fill = ln.makeelement(qn("a:solidFill"), {})
        clr = fill.makeelement(qn("a:srgbClr"), {"val": "000000"})
        fill.append(clr)
        ln.append(fill)
        tcPr.insert(k, ln)


def write_pptx_shape(ps, pages, dst: Path, tick) -> None:
    from pptx import Presentation
    from pptx.util import Emu
    prs = Presentation()
    first = ps.doc.load_page(int(pages[0])).rect if pages else None
    sw, sh = (first.width, first.height) if first is not None else (720.0, 540.0)
    prs.slide_width, prs.slide_height = Emu(_emu(sw)), Emu(_emu(sh))
    blank = prs.slide_layouts[6]
    for i, pg in enumerate(pages):
        tick(i, len(pages), f"p.{pg + 1}")
        r = ps.doc.load_page(int(pg)).rect
        scale = min(sw / r.width, sh / r.height)
        place = ((sw - r.width * scale) / 2, (sh - r.height * scale) / 2, scale)   # 크기가 다른 쪽은 줄여 가운데
        slide = prs.slides.add_slide(blank)
        edges = []
        blocks = layout.page_blocks(ps, pg, join=False, edges=edges)
        for b in blocks:                                  # 그림·표 먼저 — 글상자가 위에 오게
            if b["kind"] == "image":
                x0, y0, x1, y1 = b["rect"]
                left, top, s = place
                slide.shapes.add_picture(io.BytesIO(b["data"]), Emu(_emu(left + x0 * s)), Emu(_emu(top + y0 * s)),
                                         Emu(_emu((x1 - x0) * s)), Emu(_emu((y1 - y0) * s)))
            elif b["kind"] == "table":
                _table(slide, place, b)
        lines = [b for b in blocks if b["kind"] == "line" and b.get("rect")]
        for e in edges:                                   # 머리말·꼬리말도 그 자리에 글상자로(§1.1)
            lines.append({"kind": "line", "rect": tuple(e["rect"]), "text": e["text"], "size": e["size"],
                          "bold": e.get("bold", False), "serif": e.get("serif", True)})
        for band in layout.bands(lines):
            text, x0, width, size, bold, serif = layout.band_text(band)
            y0 = min(q["rect"][1] for q in band)
            y1 = max(q["rect"][3] for q in band)
            _text_box(slide, place, text, x0, y0, width, y1 - y0, size, bold, serif)
    tick(len(pages), len(pages), "")
    prs.save(str(dst))
