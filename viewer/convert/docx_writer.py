# -*- coding: utf-8 -*-
"""Word(편집 가능) 쓰기 — python-docx (형식 변환 SOT §1, 261014-14).

HWPX 편집 가능과 같은 블록(`layout.page_blocks(join=True)`): 이은 문단, 제목은 'Heading 1', 표는 'Table Grid' 표,
그림은 원본 크기(글 너비를 넘으면 줄인다), 쪽 사이 쪽 나눔. 쪽 크기는 첫 쪽에 맞춘다.
"""
from __future__ import annotations

import io
from pathlib import Path

from viewer.convert import layout
from viewer.convert import page_step, step_text  # 진행 창 글(형식 변환 SOT §5, 261026-2)

TITLE_STYLE = "Heading 1"      # Word 기본 서식 이름(문서 안의 스타일 키 — 화면 문구가 아니다)
TABLE_STYLE = "Table Grid"
MARGIN_PT = 56.7               # 쪽 여백 2 cm


def write_docx_edit(ps, pages, dst: Path, tick) -> None:
    from docx import Document
    from docx.enum.text import WD_BREAK
    from docx.shared import Pt
    out = Document()
    sec = out.sections[0]
    if pages:
        r = ps.doc.load_page(int(pages[0])).rect
        sec.page_width, sec.page_height = Pt(r.width), Pt(r.height)
    sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = Pt(MARGIN_PT)
    text_w = sec.page_width.pt - 2 * MARGIN_PT
    body = out.paragraphs[0] if out.paragraphs else None     # 빈 첫 문단 — 처음 쓰는 문단으로 쓴다

    def para(text="", style=None):
        nonlocal body
        if body is not None:
            p, body = body, None
            if text:
                p.add_run(text)
            if style:
                p.style = style
            return p
        return out.add_paragraph(text, style=style)

    per_page = {}
    for i, pg in enumerate(pages):
        tick(i, len(pages), page_step(ps, pages, i))
        if i:
            para().add_run().add_break(WD_BREAK.PAGE)
        for b in layout.page_blocks(ps, pg, join=True, edges=per_page.setdefault(pg, [])):
            if b["kind"] == "line":
                style = TITLE_STYLE if b.get("style") == "title" else None
                try:
                    para(b["text"], style)
                except KeyError:
                    para(b["text"])
            elif b["kind"] == "table":
                rows = b["rows"]
                t = out.add_table(rows=len(rows), cols=max(len(r) for r in rows))
                try:
                    t.style = TABLE_STYLE
                except Exception:
                    pass
                for ri, r in enumerate(rows):
                    for ci, c in enumerate(r):
                        if c:
                            t.cell(ri, ci).text = str(c)
                body = None
            else:
                x0, y0, x1, y1 = b["rect"]
                w, h = x1 - x0, y1 - y0
                if w > text_w:
                    w, h = text_w, h * text_w / w
                data = b["data"]
                try:
                    para().add_run().add_picture(io.BytesIO(data), width=Pt(w), height=Pt(h))
                except Exception:
                    continue                             # Word 가 못 받는 그림 형식 — 건너뛴다
    _header_footer(out, sec, layout.edge_plan(per_page))
    set_hf_distance(sec, per_page)
    tick(len(pages), len(pages), step_text("save"))
    out.save(str(dst))


def set_hf_distance(sec, per_page) -> None:
    """머리글·바닥글 자리 = 원본 머리말 위쪽·꼬리말 아래쪽(쪽 높이 비율로 옮긴다)."""
    from docx.shared import Pt
    ph = sec.page_height.pt
    edges = [e for es in per_page.values() for e in es]
    heads = [e["rect"][1] / float(e.get("ph") or ph) for e in edges if e["where"] == "head"]
    foots = [e["rect"][3] / float(e.get("ph") or ph) for e in edges if e["where"] == "foot"]
    if heads:
        sec.header_distance = Pt(max(0.0, min(heads) * ph))
    if foots:
        sec.footer_distance = Pt(max(0.0, (1.0 - max(foots)) * ph))


def _page_field(run) -> None:
    """런 안에 Word 쪽 번호 필드(PAGE)."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    for kind, text in (("begin", None), (None, " PAGE "), ("separate", None), (None, "1"), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        elif text == " PAGE ":
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        else:
            el = OxmlElement("w:t")
            el.text = text
        run._r.append(el)


def _header_footer(out, sec, plan) -> None:
    """머리말·꼬리말(형식 변환 SOT §1.1) — Word 머리글·바닥글 기능으로, 쪽 번호는 PAGE 필드, 홀·짝이 다르면 따로."""
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt
    al = {"LEFT": WD_ALIGN_PARAGRAPH.LEFT, "CENTER": WD_ALIGN_PARAGRAPH.CENTER, "RIGHT": WD_ALIGN_PARAGRAPH.RIGHT}
    if any(k in plan[w] for w in ("head", "foot") for k in ("ODD", "EVEN")):
        out.settings.odd_and_even_pages_header_footer = True
    for where in ("head", "foot"):
        for kind, lines in plan[where].items():
            if kind == "EVEN":
                part = sec.even_page_header if where == "head" else sec.even_page_footer
            else:
                part = sec.header if where == "head" else sec.footer
            part.is_linked_to_previous = False
            first = True
            for ln in lines:
                pw = sec.page_width.pt
                text, align, indent = layout.compose(ln, pw, sec.left_margin.pt, pw - sec.right_margin.pt)
                para = part.paragraphs[0] if first and part.paragraphs else part.add_paragraph()
                first = False
                para.alignment = al.get(align)
                pf = para.paragraph_format
                pf.left_indent = Pt(indent) if indent else None
                pf.space_before = pf.space_after = Pt(0)       # 기본 서식의 문단 뒤 간격이 바닥글을 위로 민다
                pf.line_spacing = 1.0
                bits = text.split(layout.NUM_TOKEN)
                for k, bit in enumerate(bits):
                    run = para.add_run(bit)
                    run.font.size = Pt(ln["size"])
                    run.font.bold = bool(ln["bold"])
                    run.font.name = layout.SERIF_FACE if ln["serif"] else layout.SANS_FACE
                    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), run.font.name)
                    if k < len(bits) - 1:
                        _page_field(para.add_run())
    start = plan.get("start")
    if start and int(start) > 1:
        pg = OxmlElement("w:pgNumType")
        pg.set(qn("w:start"), str(int(start)))
        sec._sectPr.append(pg)
