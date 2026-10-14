# -*- coding: utf-8 -*-
"""형식별 쓰기 (형식 변환 SOT §1·§2.3). 모두 `PageSource` 만 읽는다 — 형식마다 추출을 새로 짜지 않는다(§3).

`writer(ps, pages, dst, tick)` — `tick(했음, 전체, 글)` 은 취소면 `Cancelled` 를 던진다(부르는 쪽이 만들던 파일을 지운다).
"""
from __future__ import annotations

from pathlib import Path


def write_xlsx(ps, pages, dst: Path, tick) -> None:
    """표마다 시트 하나(`p<쪽>_<번째>`). 표가 하나도 없으면 파일을 만들지 않는다(§2.3)."""
    from openpyxl import Workbook
    from viewer.convert import NothingToWrite
    wb = Workbook()
    wb.remove(wb.active)
    n = 0
    for i, pg in enumerate(pages):
        tick(i, len(pages), f"p.{pg + 1}")
        for k, (_bbox, rows) in enumerate(ps.tables(pg), start=1):
            ws = wb.create_sheet(f"p{pg + 1}_{k}"[:31])
            for r in rows:
                ws.append([("" if c is None else str(c)) for c in r])
            n += 1
    if not n:
        raise NothingToWrite("no tables")          # 화면 문구는 실행기가 tr("표가 없습니다") 로
    tick(len(pages), len(pages), "")
    wb.save(str(dst))


def write_txt(ps, pages, dst: Path, tick) -> None:
    """본문 줄(텍스트 창과 같은 원천·고침 반영). UTF-8 BOM, 쪽 사이 빈 줄 둘(§2.3)."""
    parts = []
    for i, pg in enumerate(pages):
        tick(i, len(pages), f"p.{pg + 1}")
        parts.append("\n".join(r.get("text", "") for r in ps.lines(pg)))
    tick(len(pages), len(pages), "")
    Path(dst).write_text("\n\n\n".join(parts), encoding="utf-8-sig")


def write_png(ps, pages, dst: Path, tick) -> None:
    """같은 이름의 폴더 안에 `<이름>_p0001.png`(§2.3). 쪽 그림은 이 스레드의 문서로 그린다."""
    import fitz
    from viewer.convert import IMAGE_DPI
    dst = Path(dst)
    dst.mkdir(parents=True, exist_ok=False)
    stem = Path(ps.path).stem
    zoom = IMAGE_DPI / 72.0
    for i, pg in enumerate(pages):
        tick(i, len(pages), f"p.{pg + 1}")
        pix = ps.doc.load_page(int(pg)).get_pixmap(matrix=fitz.Matrix(zoom, zoom), alpha=False)
        pix.save(str(dst / f"{stem}_p{pg + 1:04d}.png"))
    tick(len(pages), len(pages), "")


WRITERS = {"xlsx": write_xlsx, "txt": write_txt, "png": write_png}
