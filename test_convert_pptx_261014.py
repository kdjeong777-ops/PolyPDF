# -*- coding: utf-8 -*-
"""261014-19: PowerPoint 모양 그대로 — 고칠 수 있는 글상자 + 그림 + 표 (형식 변환 SOT §1, 개선 로드맵 P3-1).

실제 처리기(`action_convert(fmt="pptx_shape")`)로:
A. 쪽마다 슬라이드, 슬라이드 크기 = 첫 쪽, 크기가 다른 쪽은 줄여 가운데
B. 원본 줄이 그 자리의 글상자(쪽 그림이 아니다) — 글자 크기·굵기·동아시아 글꼴·글자 간격
C. 그림(박힌 그림 + 선 그림)은 그림으로, 바탕 그림은 뺌, 선 그림 안 글은 그림 안에
D. 표는 진짜 표 — 칸 글, 채움 없음, 테두리 규격 차례(lnL·lnR·lnT·lnB)
E. 머리말·꼬리말도 그 자리에 글상자로
"""
import os, sys, tempfile, shutil
os.environ["QT_QPA_PLATFORM"] = "offscreen"
sys.stdout.reconfigure(encoding="utf-8")
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from pathlib import Path

import fitz
from PyQt6.QtCore import QStandardPaths
QStandardPaths.setTestModeEnabled(True)
from PyQt6.QtWidgets import QApplication, QMessageBox
from PyQt6.QtTest import QTest

app = QApplication.instance() or QApplication(sys.argv[:1])
fails = []


def chk(cond, msg, extra=""):
    print(("PASS" if cond else "FAIL"), "-", msg, extra if not cond else "")
    if not cond:
        fails.append(msg)


def spin(n=10, ms=20):
    for _ in range(n):
        app.processEvents(); QTest.qWait(ms)


root = Path(tempfile.mkdtemp(prefix="polypdf_convpptx_"))
try:
    src = root / "P.pdf"
    d = fitz.open()
    for i in range(3):
        p = d.new_page(width=595, height=842) if i < 2 else d.new_page(width=842, height=595)
        p.insert_text((400, 40), f"Guide header p.{i + 11}", fontsize=9, fontname="tiro")
        p.insert_text((72, 80), f"Slide Title {i}", fontsize=20, fontname="hebo")
        p.insert_text((90, 120), "Indented body line with some words.", fontsize=11, fontname="tiro")
        if i == 0:
            x0, y0, cw, rh = 72, 160, 140, 24
            for r in range(4):
                p.draw_line((x0, y0 + r * rh), (x0 + 3 * cw, y0 + r * rh))
            for c in range(4):
                p.draw_line((x0 + c * cw, y0), (x0 + c * cw, y0 + 3 * rh))
            for r in range(3):
                for c in range(3):
                    p.insert_text((x0 + c * cw + 6, y0 + r * rh + 16), f"R{r}C{c}", fontsize=10)
            gx, gy = 72, 340
            p.draw_rect(fitz.Rect(gx, gy, gx + 200, gy + 120))
            for k in range(6):
                p.draw_line((gx + 10 + k * 30, gy + 110), (gx + 25 + k * 30, gy + 110 - (k + 1) * 15))
            p.insert_text((gx + 80, gy + 118), "axislabel", fontsize=7)
            pm = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 60, 30), False); pm.set_rect(pm.irect, (30, 120, 200))
            p.insert_image(fitz.Rect(320, 340, 520, 440), stream=pm.tobytes("png"))
        if i == 1:
            bg = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 40, 30), False); bg.set_rect(bg.irect, (250, 250, 240))
            p.insert_image(p.rect, stream=bg.tobytes("png"), overlay=False, keep_proportion=False)
        p.insert_text((285, (815 if i < 2 else 570)), f"- {i + 11} -", fontsize=9, fontname="tiro")
    d.save(str(src)); d.close()

    QMessageBox.question = staticmethod(lambda *x, **k: QMessageBox.StandardButton.No)
    infos = []
    QMessageBox.information = staticmethod(lambda *x, **k: infos.append(x[2] if len(x) > 2 else ""))
    QMessageBox.warning = staticmethod(lambda *x, **k: infos.append(x[2] if len(x) > 2 else ""))
    from viewer.widgets import convert_dialog as cdmod
    cdmod.ConvertDialog.exec = lambda dlg: (dlg._accept(), cdmod.ConvertDialog.DialogCode.Accepted)[1]
    from viewer.app import MainWindow
    from viewer.history import HistoryItem
    mw = MainWindow(); mw.resize(1300, 850); mw.show(); app.processEvents()
    mw.open_folder(root); spin(15)
    mw._load_main(HistoryItem(str(src), 0, "", "bookmark")); spin(10)
    ok = mw.action_convert(fmt="pptx_shape"); spin(3)
    out = root / "P.pptx"
    chk(ok and out.exists(), "준비 — PPTX 가 같은 폴더·같은 이름으로", str(infos[-1:]))

    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    prs = Presentation(str(out))
    E = 12700
    chk(len(prs.slides) == 3 and abs(prs.slide_width - 595 * E) < E and abs(prs.slide_height - 842 * E) < E,
        "A 쪽마다 슬라이드, 크기 = 첫 쪽", f"{len(prs.slides)} {prs.slide_width / E} {prs.slide_height / E}")

    def boxes(sl):
        return [s for s in sl.shapes if s.has_text_frame and s.shape_type == MSO_SHAPE_TYPE.TEXT_BOX]

    s0 = prs.slides[0]
    tb = {s.text_frame.text: s for s in boxes(s0)}
    t = tb.get("Slide Title 0")
    ind = tb.get("Indented body line with some words.")
    chk(t is not None and ind is not None and abs(ind.left / E - 90) < 1.5 and abs(t.left / E - 72) < 1.5,
        "B 원본 줄이 그 자리의 글상자(들여쓰기 포함)", str(list(tb)[:6]))
    if t is not None:
        r = t.text_frame.paragraphs[0].runs[0]
        rpr = r._r.get_or_add_rPr().xml
        chk(abs(r.font.size.pt - 20) < 0.6 and r.font.bold and 'a:ea typeface="맑은 고딕"' in rpr,
            "B 글자 크기·굵기·동아시아 글꼴 = 원본", rpr[:200])
    if ind is not None:
        rpr = ind.text_frame.paragraphs[0].runs[0]._r.get_or_add_rPr()
        from viewer.convert import layout
        if layout.measure("ab", 10, True) is not None:
            chk(rpr.get("spc") not in (None, "0"), "B 글자 간격을 원본 줄 너비에 맞춘다", str(rpr.get("spc")))
    pics = [s for s in s0.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE]
    chk(len(pics) == 2 and "axislabel" not in tb, "C 그림 둘(선 그림 + 박힌 그림), 선 그림 안 글은 그림 안에", str(len(pics)))
    s1 = prs.slides[1]
    chk(not [s for s in s1.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE], "C 바탕 그림은 뺀다")
    tables = [s for s in s0.shapes if s.has_table]
    if tables:
        tblx = tables[0].table
        print("  (표 크기)", len(tblx.rows), len(tblx.columns), [[c.text for c in r.cells] for r in tblx.rows])
        cells = [tblx.cell(r, c).text for r in range(min(3, len(tblx.rows))) for c in range(min(3, len(tblx.columns)))]
        xml = tblx.cell(0, 0)._tc.get_or_add_tcPr().xml
        order = [xml.find(k) for k in ("<a:lnL", "<a:lnR", "<a:lnT", "<a:lnB")]
        chk(cells == [f"R{r}C{c}" for r in range(3) for c in range(3)] and "a:noFill" in xml and order == sorted(order) and -1 not in order,
            "D 진짜 표 — 칸 글·채움 없음·테두리 규격 차례", str(cells[:3]))
    else:
        chk(False, "D 진짜 표")
    chk(not any(x.startswith("R1C1") for x in tb), "D 표 글이 글상자로 겹쳐 나오지 않는다")
    chk("Guide header p.11" in tb and "- 11 -" in tb, "E 머리말·꼬리말도 그 자리에 글상자로", str(list(tb)))
    s2 = prs.slides[2]
    t2 = {s.text_frame.text: s for s in boxes(s2)}.get("Slide Title 2")
    sc = min(595 / 842, 842 / 595)
    chk(t2 is not None and abs(t2.left / E - 72 * sc) < 2 and t2.top / E > 100,
        "A 크기가 다른(가로) 쪽은 비율대로 줄여 가운데", f"{t2.left / E if t2 else None} {t2.top / E if t2 else None}")
finally:
    shutil.rmtree(root, ignore_errors=True)

print("\n=== " + ("ALL PASS" if not fails else f"FAILURE ({len(fails)})") + " ===")
for m in fails:
    print(" -", m)
sys.stdout.flush()
os._exit(1 if fails else 0)
