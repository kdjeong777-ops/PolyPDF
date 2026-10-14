# -*- coding: utf-8 -*-
"""261014-13: Word(편집 가능) — 가져와 고친 pdf2docx 를 보조 프로세스에서 (형식 변환 SOT §4.1, 개선 로드맵 P2-3).

실제 처리기(`action_convert(fmt="docx_edit")`)로:
A. docx 에 제목·문단 글이 들어간다
B. 선 그은 표가 Word **표**로 나온다(칸 글 그대로)
C. 그림이 들어간다 — 회전해 넣은 그림도(고침 2: opencv 대신 Pillow)
D. 스캔 쪽(글자층 없음)은 우리 OCR 결과가 글로 들어간다(임시 사본 — 원본 옆·임시 폴더에 흔적이 남지 않는다)
E. 변환은 **보조 프로세스**에서 — 이 프로세스에는 pdf2docx 가 실리지 않는다(opencv 도)
F. 취소하면 docx 를 지운다(고침 3)
"""
import os, sys, tempfile, shutil
os.environ["QT_QPA_PLATFORM"] = "offscreen"
sys.stdout.reconfigure(encoding="utf-8")
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from pathlib import Path

import fitz
from PyQt6.QtCore import QStandardPaths
QStandardPaths.setTestModeEnabled(True)
from PyQt6.QtWidgets import QApplication, QMessageBox
from PyQt6.QtTest import QTest

app = QApplication.instance() or QApplication(sys.argv[:1])
fails = []


def chk(cond, msg, extra=""):
    print(("PASS" if cond else "FAIL"), "-", msg, extra if not cond else "")
    if not cond:
        fails.append(msg)


def spin(n=10, ms=20):
    for _ in range(n):
        app.processEvents(); QTest.qWait(ms)


root = Path(tempfile.mkdtemp(prefix="polypdf_convdocx_"))
try:
    src = root / "W.pdf"
    d = fitz.open()
    p1 = d.new_page(width=595, height=842)
    p1.insert_text((72, 80), "Asphalt Mixture Design", fontsize=18)
    for k in range(3):
        p1.insert_text((72, 120 + k * 16), f"Paragraph line {k} explains the compaction method.", fontsize=11)
    x0, y0, cw, rh = 72, 220, 140, 24
    for r in range(4):
        p1.draw_line((x0, y0 + r * rh), (x0 + 3 * cw, y0 + r * rh))
    for c in range(4):
        p1.draw_line((x0 + c * cw, y0), (x0 + c * cw, y0 + 3 * rh))
    for r in range(3):
        for c in range(3):
            p1.insert_text((x0 + c * cw + 6, y0 + r * rh + 16), f"R{r}C{c}", fontsize=10)
    pm = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 60, 30), False); pm.set_rect(pm.irect, (200, 30, 30))
    png = pm.tobytes("png")
    p2 = d.new_page(width=595, height=842)
    p2.insert_text((72, 80), "Figure page", fontsize=14)
    p2.insert_image(fitz.Rect(72, 120, 272, 220), stream=png)
    p2.insert_image(fitz.Rect(300, 120, 400, 320), stream=png, rotate=90)
    p3 = d.new_page(width=595, height=842)                       # 스캔 쪽 — 글자층 없이 그림만
    scan = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 200, 280), False); scan.set_rect(scan.irect, (250, 250, 250))
    p3.insert_image(p3.rect, stream=scan.tobytes("png"))
    d.save(str(src)); d.close()

    # 3쪽의 우리 OCR 결과(단어장 SOT 소유 study.db) — 실제 OCR 대신 결과를 넣어 둔다
    from viewer.study.study_store import StudyStore, file_key_for
    st = StudyStore()
    st.save_page(file_key_for(str(src)), 2, "SCANNED OCR SENTENCE", dpi=300, engine="tesseract", source="ocr",
                 conf=90.0, lang="eng",
                 words=[{"surface": w, "x0": 300 + i * 260, "y0": 900, "x1": 540 + i * 260, "y1": 960}
                        for i, w in enumerate(["SCANNED", "OCR", "SENTENCE"])])
    st.close()

    QMessageBox.question = staticmethod(lambda *x, **k: QMessageBox.StandardButton.No)
    infos = []
    QMessageBox.information = staticmethod(lambda *x, **k: infos.append(x[2] if len(x) > 2 else ""))
    QMessageBox.warning = staticmethod(lambda *x, **k: infos.append(x[2] if len(x) > 2 else ""))
    from viewer.widgets import convert_dialog as cdmod
    cdmod.ConvertDialog.exec = lambda dlg: (dlg._accept(), cdmod.ConvertDialog.DialogCode.Accepted)[1]

    from viewer import index_proc
    calls = []
    _real_run = index_proc.run_job

    def spy_run(job, *a, **k):
        calls.append(k.get("server_arg"))
        return _real_run(job, *a, **k)
    index_proc.run_job = spy_run

    from viewer.app import MainWindow
    from viewer.history import HistoryItem
    mw = MainWindow(); mw.resize(1300, 850); mw.show(); app.processEvents()
    mw.open_folder(root); spin(15)
    mw._load_main(HistoryItem(str(src), 0, "", "bookmark")); spin(10)
    tmp_before = {p.name for p in Path(tempfile.gettempdir()).glob("polypdf_conv_*")}
    ok = mw.action_convert(fmt="docx_edit"); spin(3)
    out = root / "W.docx"
    chk(ok and out.exists(), "준비 — Word 파일이 같은 폴더·같은 이름으로", str(infos[-1:]))
    if out.exists():
        from docx import Document
        doc = Document(str(out))
        body = "\n".join(p.text for p in doc.paragraphs)
        cells = [c.text.strip() for t in doc.tables for row in t.rows for c in row.cells]
        chk("Asphalt Mixture Design" in body and "Paragraph line 2" in body, "A 제목·문단 글", body[:120])
        chk(len(doc.tables) >= 1 and "R1C1" in cells and "R2C2" in cells, "B 선 그은 표가 Word 표로", str(cells[:9]))
        n_img = len(doc.inline_shapes) + sum(1 for r in doc.element.body.iter() if r.tag.endswith("}anchor"))
        chk(n_img >= 2, "C 그림이 들어간다(회전해 넣은 그림 포함)", f"images={n_img}")
        chk("SCANNED OCR SENTENCE" in body or "SCANNED" in body, "D 스캔 쪽은 우리 OCR 결과가 글로", body[-200:])
    left = {p.name for p in Path(tempfile.gettempdir()).glob("polypdf_conv_*")} - tmp_before
    chk(not left and not list(root.glob("*_textlayer_tmp.pdf")), "D 임시 사본을 남기지 않는다(원본 옆·임시 폴더)", str(left))
    chk(calls.count("--convert-server") == 1, "E 변환은 보조 프로세스(--convert-server)에서(나머지는 같은 때 돈 색인)", str(calls))
    chk("viewer._vendor.pdf2docx" not in sys.modules and "cv2" not in sys.modules,
        "E 이 프로세스에는 pdf2docx·opencv 가 실리지 않는다")

    # ── F. 취소 ──────────────────────────────────────────────────
    from viewer import convert
    seen = []

    def cancel_soon(i, n, label):
        seen.append(label)
        return len(seen) < 2                     # 자식이 첫 진행을 보낸 뒤 취소
    try:
        convert.run(str(src), "docx_edit", pages=[0, 1], progress=cancel_soon)
        cancelled = False
    except convert.Cancelled:
        cancelled = True
    chk(cancelled and not (root / "W (2).docx").exists(), "F 취소하면 docx 를 지운다", str(seen[:3]))
finally:
    index_proc.run_job = _real_run if "_real_run" in dir() else index_proc.run_job
    shutil.rmtree(root, ignore_errors=True)

print("\n=== " + ("ALL PASS" if not fails else f"FAILURE ({len(fails)})") + " ===")
for m in fails:
    print(" -", m)
sys.stdout.flush()
os._exit(1 if fails else 0)
