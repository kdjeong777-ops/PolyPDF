# -*- coding: utf-8 -*-
"""261026-2: 변환 진행 창 글·형식 차례 (형식 변환 SOT §2.1·§5, 디자인 SOT §2.7 '진행 창 글').

사용자 보고: 전체 변환을 취소하고 1-50쪽으로 다시 변환하자 진행 창에 `24 ….pdf  parse 33  (320/1000)  |||` —
원인은 취소 뒤 남은 상태가 아니라 **엔진 내부 값**(pdf2docx 단계 이름 `parse`, 파일마다 나눈 척도 1000)이 그대로 보인 것.
가운데 정렬 글은 하트비트(`|`→`||`→`|||`)가 폭을 바꿀 때마다 출렁거렸다.

실제 처리기(`action_convert` → `_run_merge_job` 진행 창)로:
A. 형식 차례 — 형식마다 '모양 그대로' 가 위, 'Word (Microsoft Word 사용)' 은 'Word (모양 그대로)' 바로 아래, 변환 창의 처음 고른 형식 = 'Word (모양 그대로)'
B. Word(모양 그대로) 전체 변환을 진행 창 [취소] 로 멈춘 뒤 쪽 범위 11-20 으로 다시 변환하면 결과에 그 쪽만, 진행 창 글은 `단계 k/10 · 원본 p쪽`
C. 모든 형식의 진행 글이 `단계 k/n · 원본 p쪽` 꼴 — 날 글(`p.12`·`parse 33`·`make 12`)·내부 척도(/1000)가 없다, 범위 변환은 그 쪽만(편집 가능·HWPX·텍스트)
D. 실제 진행 창 글 폭이 하트비트·쪽 수가 바뀌어도 같다(`counts=False` 변환 꼴과 `(k/n)` 을 붙이는 다른 일 꼴 둘 다)
수정 전 코드로 A·B·C·D 모두 실패를 확인했다.
"""
import os, sys, re, tempfile, shutil, time
os.environ["QT_QPA_PLATFORM"] = "offscreen"
sys.stdout.reconfigure(encoding="utf-8")
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from pathlib import Path

import fitz
from PyQt6.QtCore import QStandardPaths, QTimer
QStandardPaths.setTestModeEnabled(True)
from PyQt6.QtWidgets import QApplication, QProgressDialog, QPushButton, QLabel
from PyQt6.QtTest import QTest

app = QApplication.instance() or QApplication(sys.argv[:1])
import modal_guard
# 진행 창(창 모달 QProgressDialog)은 오래 떠 있는 것이 정상이라 감시 타이머는 끈다(켜면 진행 창을 닫아 변환이 취소된다)
guard = modal_guard.install(label=lambda: "변환 진행", watch_ms=0, echo=False)
fails = []


def chk(cond, msg, extra=""):
    print(("PASS" if cond else "FAIL"), "-", msg, extra if not cond else "")
    if not cond:
        fails.append(msg)


def spin(n=10, ms=20):
    for _ in range(n):
        app.processEvents(); QTest.qWait(ms)


# 진행 글 2줄 꼴: 단계 k/n · 원본 p쪽 (숫자 앞 채움은 U+2007)
STEP = re.compile(r"^(쪽 변환|쪽 분석|Word 쓰기|OCR 글자층|머리말·꼬리말 찾기|저장|Word 로 변환)"
                  r"( [ \d]+/\d+)?(  ·  원본 [ \d]+쪽)?$")
RAW = re.compile(r"\bparse\b|\bmake\b|\bp\.\d|/1000\b")

root = Path(tempfile.mkdtemp(prefix="polypdf_convprog_"))
try:
    src = root / "R.pdf"
    d = fitz.open()
    for i in range(60):
        p = d.new_page()
        p.insert_text((72, 100), f"PAGEMARK{i + 1:03d} body text line", fontsize=12)
    d.save(str(src)); d.close()

    # ── A. 형식 차례·처음 고른 형식 ─────────────────────────────────
    from viewer import convert
    keys = [f.key for f in convert.FORMATS]
    chk(keys.index("docx_shape") < keys.index("docx_edit") and keys.index("hwpx_shape") < keys.index("hwpx_edit"),
        "A 형식마다 '모양 그대로' 가 '편집 가능' 위", str(keys))
    chk(keys.index("docx_word") == keys.index("docx_shape") + 1, "A 'Word (Microsoft Word 사용)' 은 'Word (모양 그대로)' 바로 아래", str(keys))
    from viewer.widgets.convert_dialog import ConvertDialog
    dlg0 = ConvertDialog(None, files=[str(src)], page_count=60)
    chk(dlg0.format_key() == "docx_shape", "A 변환 창의 처음 고른 형식 = Word (모양 그대로)", str(dlg0.format_key()))
    dlg0.deleteLater()

    # 변환 창: 쪽 범위를 정해 '변환' 을 누른 것과 같게(실제 _accept)
    want = {"range": None}

    def fake_exec(dlg):
        if want["range"]:
            dlg.rb_range.setChecked(True); dlg.ed_range.setText(want["range"])
        dlg._accept()
        return ConvertDialog.DialogCode.Accepted if dlg.result() else ConvertDialog.DialogCode.Rejected
    ConvertDialog.exec = fake_exec

    from viewer.app import MainWindow
    from viewer.history import HistoryItem
    mw = MainWindow(); mw.resize(1300, 850); mw.show(); app.processEvents()
    mw.open_folder(root); spin(15)
    mw._load_main(HistoryItem(str(src), 0, "", "bookmark")); spin(10)

    # 진행 창 글을 20ms 마다 받아 적고, 조건이 되면 진행 창의 [취소] 단추를 누른다
    shot = {"labels": [], "cancel_when": None, "widths": []}

    def sample():
        for pd in mw.findChildren(QProgressDialog):
            if not pd.isVisible() and not pd.labelText():
                continue
            t = pd.labelText()
            if not shot["labels"] or shot["labels"][-1] != t:
                shot["labels"].append(t)
            lab = pd.findChild(QLabel)
            if lab is not None:
                fm = lab.fontMetrics()
                shot["widths"].append((t, tuple(fm.horizontalAdvance(x) for x in t.split("\n"))))
            cw = shot["cancel_when"]
            if cw and cw(t):
                shot["cancel_when"] = None
                btn = pd.findChild(QPushButton)
                if btn is not None:
                    btn.click()
    sampler = QTimer(); sampler.setInterval(20); sampler.timeout.connect(sample); sampler.start()

    def line2(t):
        """진행 글의 단계 줄(마지막 줄)에서 하트비트를 뗀다."""
        ln = t.split("\n")[-1]
        return re.sub(r"\s+\S{3}$", "", ln)

    # ── B. 전체 변환 → 진행 창 [취소] → 11-20 으로 다시 ─────────────────
    shot.update(labels=[], cancel_when=lambda t: "쪽 분석" in t or "parse" in t)
    want["range"] = None
    ok1 = mw.action_convert(fmt="docx_shape"); spin(5)
    chk(not ok1 and not (root / "R.docx").exists(), "B 준비 — 전체 변환을 진행 창 [취소] 로 멈추면 결과가 없다",
        str(shot["labels"][-3:]))
    shot.update(labels=[], cancel_when=None)
    want["range"] = "11-20"
    ok2 = mw.action_convert(fmt="docx_shape"); spin(5)
    out = root / "R.docx"
    marks = []
    if out.exists():
        from docx import Document
        marks = sorted(set(re.findall(r"PAGEMARK(\d{3})", "\n".join(p.text for p in Document(str(out)).paragraphs))))
    chk(ok2 and marks == ["%03d" % k for k in range(11, 21)], "B 취소 뒤 쪽 범위 11-20 변환 — 결과에 그 쪽만", str(marks))
    labs = [t for t in shot["labels"] if "\n" in t]
    chk(labs and all(t.split("\n")[0] == "R.pdf" for t in labs), "B 진행 글 1줄 = 파일 이름", str(shot["labels"][:3]))
    steps = [line2(t) for t in labs]
    chk(steps and all(STEP.match(s) for s in steps) and not any(RAW.search(t) for t in shot["labels"]),
        "B 진행 글 2줄 = '단계 k/n · 원본 p쪽' — parse·(…/1000) 같은 내부 값 없음",
        str([s for s in steps if not STEP.match(s)][:3] or shot["labels"][:3]))
    chk(any(s.startswith("쪽 분석") and "/10  ·" in s for s in steps) and any(s.startswith("Word 쓰기") for s in steps),
        "B 범위 변환의 쪽 수는 범위(…/10) — 쪽 분석·Word 쓰기", str(steps[:6]))
    if out.exists():
        out.unlink()

    # ── C. 모든 형식의 진행 글·범위 ─────────────────────────────────
    for key in ("docx_edit", "hwpx_edit", "hwpx_shape", "pptx_shape", "txt", "png", "xlsx", "docx_shape"):
        seen = []
        try:
            res = convert.run(str(src), key, list(range(10, 20)), progress=lambda i, n, l: seen.append(l) or True)
        except convert.NothingToWrite:
            res = None
        bad = [l for l in seen if not STEP.match(l)]
        chk(seen and not bad, f"C {key} 진행 글이 모두 '단계 k/n · 원본 p쪽' 꼴", str(bad[:3]))
        chk(any("/10  ·  원본 11쪽" in l for l in seen), f"C {key} 범위의 쪽 수(/10)·원본 쪽 번호", str(seen[:3]))
        if res is not None and key in ("docx_edit", "txt", "hwpx_edit"):
            if key == "txt":
                txt = res.read_text(encoding="utf-8-sig")
            elif key == "docx_edit":
                from docx import Document
                txt = "\n".join(p.text for p in Document(str(res)).paragraphs)
            else:
                import zipfile
                with zipfile.ZipFile(str(res)) as z:
                    txt = " ".join(z.read(n).decode("utf-8", "ignore") for n in z.namelist() if n.endswith(".xml"))
            mk = sorted(set(re.findall(r"PAGEMARK(\d{3})", txt)))
            chk(mk == ["%03d" % k for k in range(11, 21)], f"C {key} 범위 변환 — 그 쪽만", str(mk))
        if res is not None:
            shutil.rmtree(res, ignore_errors=True) if Path(res).is_dir() else Path(res).unlink(missing_ok=True)

    # ── D. 실제 진행 창 글 폭 — 하트비트·쪽 수가 바뀌어도 같다 ───────────────
    def slow_job(fmt_label):
        def job(progress):
            for k in range(1, 61):
                progress(k, 60, fmt_label(k))
                time.sleep(0.03)
        return job
    for name, counts, lab in (
            ("변환 꼴(counts=False)", False,          # 수정 전 코드에는 step_text 가 없다 — 그때의 날 글로
             lambda k: "R.pdf\n" + (convert.step_text("parse", k, 60, k, 60) if hasattr(convert, "step_text")
                                    else "parse %d" % k)),
            ("다른 일 꼴((k/n) 붙임)", True, lambda k: "병합 중")):
        shot.update(labels=[], widths=[], cancel_when=None)
        import inspect                                   # 수정 전 코드에는 counts 가 없다
        kw = {"counts": counts} if "counts" in inspect.signature(mw._run_merge_job).parameters else {}
        mw._run_merge_job(slow_job(lab), "변환", **kw); spin(2)
        mine = [(t, w) for t, w in shot["widths"] if t.startswith("R.pdf" if not counts else "병합 중")]   # '준비 중…' 은 뺀다
        last = {w[-1] for _t, w in mine}
        chk(len({t for t, _w in mine}) >= 8 and len(last) == 1,
            f"D {name} — 진행 글 폭이 하트비트·쪽 수가 바뀌어도 같다", f"widths={sorted(last)[:6]} n={len(mine)}")
    sampler.stop()
finally:
    shutil.rmtree(root, ignore_errors=True)

print("\n=== " + ("ALL PASS" if not fails else f"FAILURE ({len(fails)})") + " ===")
for m in fails:
    print(" -", m)
sys.stdout.flush()
os._exit(1 if fails else 0)
