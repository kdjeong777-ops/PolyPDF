# -*- coding: utf-8 -*-
"""261014-14: Word(편집 가능)·한글 HWPX(편집 가능·모양 그대로) (형식 변환 SOT §1, 개선 로드맵 P2-4·P2-5).

두 방식 모두 **고칠 수 있는 글**이다 — 쪽 그림 한 장이 아니다(사용자 지시 261014-14).
A. Word 편집 가능 — 제목 'Heading 1', 문단, 표(칸 글), 그림(박힌 그림 + 선 그림), 그림 제목은 글, 선 그림 안 글은 그림 안에, 쪽 나눔
B. HWPX 편집 가능(실제 처리기 `action_convert`) — 개요 1, 문단, 표, 그림, 위에서 아래 차례, 쪽 나눔, 표 자리 표시 줄이 글로 새지 않는다
C. HWPX 모양 그대로 — 쪽마다 구역(쪽 크기 = PDF 쪽), 원본 줄마다 문단, 왼쪽 여백 = 줄의 x, 위 간격, 글자 크기·굵기, 자간,
   나란한 줄은 한 문단에 빈칸으로, 그림·표는 종이 기준 원본 자리에 띄움, 바탕 그림은 뺌
D. 만든 HWPX 를 python-hwpx 로 다시 열 수 있다(한글에서 여는 확인은 사용자 PC — 형식 변환 SOT §7)
E. 머리말·꼬리말(261014-15, §1.1) — Word·HWPX 의 머리말·꼬리말 기능으로, 쪽 번호는 필드(원본 번호부터), 자리는 원본대로,
   본문에는 남지 않는다. Word 모양 그대로(pdf2docx)도.
"""
import os, sys, tempfile, shutil, zipfile, re
os.environ["QT_QPA_PLATFORM"] = "offscreen"
sys.stdout.reconfigure(encoding="utf-8")
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from pathlib import Path

import fitz
from PyQt6.QtCore import QStandardPaths
QStandardPaths.setTestModeEnabled(True)
from PyQt6.QtWidgets import QApplication, QMessageBox
from PyQt6.QtTest import QTest

app = QApplication.instance() or QApplication(sys.argv[:1])
fails = []


def chk(cond, msg, extra=""):
    print(("PASS" if cond else "FAIL"), "-", msg, extra if not cond else "")
    if not cond:
        fails.append(msg)


def spin(n=10, ms=20):
    for _ in range(n):
        app.processEvents(); QTest.qWait(ms)


def section_xml(path):
    z = zipfile.ZipFile(str(path))
    names = sorted(n for n in z.namelist() if re.match(r"Contents/section\d+\.xml$", n))
    return [z.read(n).decode("utf-8") for n in names]


def make_pdf(path):
    d = fitz.open()
    p1 = d.new_page(width=595, height=842)
    p1.insert_text((72, 80), "Mixture Design Title", fontsize=20, fontname="hebo")
    p1.insert_text((72, 120), "Body text before the table.", fontsize=11, fontname="tiro")
    p1.insert_text((90, 136), "Indented second line.", fontsize=11, fontname="tiro")
    x0, y0, cw, rh = 72, 160, 140, 24
    for r in range(4):
        p1.draw_line((x0, y0 + r * rh), (x0 + 3 * cw, y0 + r * rh))
    for c in range(4):
        p1.draw_line((x0 + c * cw, y0), (x0 + c * cw, y0 + 3 * rh))
    for r in range(3):
        for c in range(3):
            p1.insert_text((x0 + c * cw + 6, y0 + r * rh + 16), f"R{r}C{c}", fontsize=10)
    p1.insert_text((72, 300), "Body text after the table.", fontsize=11, fontname="tiro")
    gx, gy = 72, 340                                       # 선으로 그린 그림 + 그 안의 축 글 + 그림 제목
    p1.draw_rect(fitz.Rect(gx, gy, gx + 200, gy + 120))
    for k in range(6):
        p1.draw_line((gx + 10 + k * 30, gy + 110), (gx + 25 + k * 30, gy + 110 - (k + 1) * 15))
    p1.insert_text((gx + 80, gy + 118), "axislabel", fontsize=7)
    p1.insert_text((gx, gy + 140), "Figure 1. Vector chart caption", fontsize=10, fontname="tiro")
    p1.insert_text((72, 520), "Left column", fontsize=11, fontname="tiro")    # 같은 높이 두 조각
    p1.insert_text((320, 520), "Right column", fontsize=11, fontname="tiro")
    p2 = d.new_page(width=842, height=595)                 # 가로 쪽 — 쪽 크기가 따라가는지
    bg = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 40, 30), False); bg.set_rect(bg.irect, (250, 250, 240))
    p2.insert_image(p2.rect, stream=bg.tobytes("png"))     # 바탕 그림 — 넣지 않는다
    p2.insert_text((72, 80), "Landscape page text", fontsize=11)
    pm = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 60, 30), False); pm.set_rect(pm.irect, (30, 120, 200))
    p2.insert_image(fitz.Rect(72, 120, 272, 220), stream=pm.tobytes("png"))
    p2.insert_text((72, 240), "Figure 2. Raster image caption", fontsize=10)
    d.save(str(path)); d.close()


root = Path(tempfile.mkdtemp(prefix="polypdf_convhwpx_"))
try:
    src = root / "H.pdf"
    make_pdf(src)
    from viewer import convert

    # ── A. Word 편집 가능 ──────────────────────────────────────────
    out_a = convert.run(str(src), "docx_edit")
    from docx import Document
    dd = Document(str(out_a))
    paras = [(p.style.name, p.text) for p in dd.paragraphs]
    texts = [t for _s, t in paras]
    chk(("Heading 1", "Mixture Design Title") in paras, "A 제목은 Heading 1", str(paras[:3]))
    chk("Body text before the table." in texts and "Landscape page text" in texts, "A 문단 글")
    cells = [c.text for t in dd.tables for row in t.rows for c in row.cells]
    chk(len(dd.tables) == 1 and "R1C1" in cells and "R2C2" in cells, "A 표는 Word 표로", str(cells[:9]))
    chk(not any("R1C1" in t or "[표" in t for t in texts), "A 표 글·표 자리 표시가 문단으로 새지 않는다")
    chk(len(dd.inline_shapes) == 2, "A 그림 둘(선 그림 + 박힌 그림, 바탕 그림은 뺌)", str(len(dd.inline_shapes)))
    chk("Figure 1. Vector chart caption" in texts and "Figure 2. Raster image caption" in texts, "A 그림 제목은 글")
    chk(not any("axislabel" in t for t in texts), "A 선 그림 안의 글은 그림 안에(글로 따로 넣지 않는다)")
    chk(any('w:type="page"' in p._p.xml for p in dd.paragraphs), "A 쪽 사이 쪽 나눔")

    # ── B. HWPX 편집 가능(실제 처리기) ───────────────────────────────
    QMessageBox.question = staticmethod(lambda *x, **k: QMessageBox.StandardButton.No)
    infos = []
    QMessageBox.information = staticmethod(lambda *x, **k: infos.append(x[2] if len(x) > 2 else ""))
    QMessageBox.warning = staticmethod(lambda *x, **k: infos.append(x[2] if len(x) > 2 else ""))
    from viewer.widgets import convert_dialog as cdmod
    cdmod.ConvertDialog.exec = lambda dlg: (dlg._accept(), cdmod.ConvertDialog.DialogCode.Accepted)[1]
    from viewer.app import MainWindow
    from viewer.history import HistoryItem
    mw = MainWindow(); mw.resize(1300, 850); mw.show(); app.processEvents()
    mw.open_folder(root); spin(15)
    mw._load_main(HistoryItem(str(src), 0, "", "bookmark")); spin(10)
    ok = mw.action_convert(fmt="hwpx_edit"); spin(3)
    out_b = root / "H.hwpx"
    chk(ok and out_b.exists(), "B HWPX 가 같은 폴더·같은 이름으로", str(infos[-1:]))
    if out_b.exists():
        x = "".join(section_xml(out_b))
        i_title, i_tbl, i_after = x.find("Mixture Design Title"), x.find("<hp:tbl"), x.find("Body text after the table.")
        chk(i_title >= 0 and "Body text before the table." in x, "B 제목·문단 글")
        chk(x.count("<hp:tbl") == 1 and all(f"R{r}C{c}" in x for r in range(3) for c in range(3)), "B 표 칸 글",
            str(x.count("<hp:tbl")))
        chk("R1C1" not in x[:i_tbl] if i_tbl >= 0 else False, "B 표 글이 문단으로 겹쳐 나오지 않는다(표 글은 표로만)")
        chk("[표" not in x, "B 표 자리 표시 줄이 글로 새지 않는다")
        chk(0 <= i_title < i_tbl < i_after, "B 위에서 아래 차례(제목 → 표 → 표 뒤 문단)", f"{i_title} {i_tbl} {i_after}")
        chk(x.count("<hp:pic") == 2, "B 그림 둘(바탕 그림은 뺌)", str(x.count("<hp:pic")))
        chk("Figure 1. Vector chart caption" in x and "axislabel" not in x, "B 그림 제목은 글, 선 그림 안 글은 그림 안에")
        chk(x.count('pageBreak="1"') >= 1, "B 쪽 사이 쪽 나눔")
        from hwpx import HwpxDocument
        try:
            HwpxDocument.open(str(out_b)); reopened = True
        except Exception:
            reopened = False
        chk(reopened, "D 만든 HWPX 를 다시 열 수 있다")

    # ── C. HWPX 모양 그대로 ────────────────────────────────────────
    out_c = convert.run(str(src), "hwpx_shape")
    xs = section_xml(out_c)
    full = "".join(xs)
    z = zipfile.ZipFile(str(out_c))
    hdr = z.read("Contents/header.xml").decode("utf-8")
    sizes = re.findall(r'<hp:pagePr[^>]*width="(\d+)"[^>]*height="(\d+)"', full)
    chk(len(xs) == 2 and sizes == [("59500", "84200"), ("84200", "59500")],
        "C 쪽마다 구역, 쪽 크기 = PDF 쪽(HWPUNIT)", str(sizes))
    chk("<hp:pic" in full and z.namelist().count("BinData/BIN0001.png") == 1
        and len([n for n in z.namelist() if n.startswith("BinData/")]) == 2,
        "C 그림 둘(선 그림 + 박힌 그림) — 바탕 그림은 뺀다, 쪽 그림 한 장이 아니다",
        str([n for n in z.namelist() if n.startswith("BinData/")]))

    def para_of(text):
        for m in re.finditer(r'<hp:p [^>]*>(?:(?!<hp:p ).)*?</hp:p>', full, re.S):
            if text in m.group(0):
                return m.group(0)
        return ""

    def para_pr(pid):
        m = re.search(r'<hh:paraPr [^>]*id="%s"[^>]*>.*?</hh:paraPr>' % pid, hdr, re.S)
        return m.group(0) if m else ""

    def char_pr(cid):
        m = re.search(r'<hh:charPr [^>]*id="%s"[^>]*>.*?</hh:charPr>' % cid, hdr, re.S)
        return m.group(0) if m else ""

    def margin(pp, name):
        m = re.search(r'<hp:case[^>]*>.*?<hc:%s value="(-?\d+)"' % name, pp, re.S)
        return int(m.group(1)) if m else None

    lines = ["Mixture Design Title", "Body text before the table.", "Indented second line.",
             "Body text after the table.", "Figure 1. Vector chart caption", "Landscape page text"]
    ps_ = {t: para_of(t) for t in lines}
    chk(all(ps_[t] and len(re.findall(r"<hp:t>", ps_[t])) == 1 for t in lines), "C 원본 줄마다 문단 하나(글로 — 고칠 수 있다)",
        str([t for t in lines if not ps_[t]]))
    pp = {t: para_pr(re.search(r'paraPrIDRef="(\d+)"', ps_[t]).group(1)) for t in lines if ps_[t]}
    cp = {t: char_pr(re.search(r'charPrIDRef="(\d+)"', ps_[t]).group(1)) for t in lines if ps_[t]}
    page_left = int(re.search(r'<hp:margin [^>]*left="(\d+)"', xs[0]).group(1))       # 쪽 왼쪽 여백 = 원본 글 영역 왼쪽
    xi = page_left + (margin(pp.get("Indented second line.", ""), "left") or 0)
    xb = page_left + (margin(pp.get("Body text before the table.", ""), "left") or 0)
    chk(abs(xi - 9000) <= 60 and abs(xb - 7200) <= 60 and xi - xb == 1800,
        "C 쪽 왼쪽 여백 + 문단 왼쪽 여백 = 원본 줄의 x(들여쓰기)", f"{page_left} {xi} {xb}")
    t_prev = margin(pp.get("Mixture Design Title", ""), "prev") or 0
    a_prev = margin(pp.get("Body text after the table.", ""), "prev") or 0
    chk(5000 < t_prev < 7000 and a_prev > 4000, "C 위 간격 = 원본 줄 자리(표가 띄워져도 표 높이만큼 비운다)", f"{t_prev} {a_prev}")
    chk('height="2000"' in cp.get("Mixture Design Title", "") and "<hh:bold" in cp.get("Mixture Design Title", "")
        and 'height="1100"' in cp.get("Body text before the table.", "")
        and "<hh:bold" not in cp.get("Body text before the table.", ""),
        "C 글자 크기·굵기 = 원본")
    from viewer.convert import layout
    if layout.measure("ab", 10, True) is not None:
        sp = [re.search(r'<hh:spacing hangul="(-?\d+)"', cp[t]).group(1) for t in cp]
        chk(any(v != "0" for v in sp), "C 자간을 원본 줄 너비에 맞춘다", str(sp))
    lr = para_of("Left column")
    chk(lr and "Right column" in lr and re.search(r"Left column {10,}Right column", lr),
        "C 나란한 줄은 한 문단에 빈칸으로 벌린다", lr[-200:])
    poss = re.findall(r'<hp:pos [^>]*treatAsChar="0"[^>]*>', full)
    chk(len(poss) == 3 and all('vertRelTo="PAPER"' in p and 'horzRelTo="PAPER"' in p for p in poss),
        "C 그림 둘·표 하나를 종이 기준 자리에 띄운다", str(len(poss)))
    tbl = re.search(r'<hp:tbl [^>]*>.*?<hp:pos [^>]*>', full, re.S)
    chk(tbl and 'vertOffset="16000"' in tbl.group(0) and 'horzOffset="7200"' in tbl.group(0)
        and "R2C2" in full, "C 표는 원본 자리(72, 160 pt)에 진짜 표로")
    chk("Figure 1. Vector chart caption" in full and "axislabel" not in full and "[표" not in full,
        "C 그림 제목은 글, 선 그림 안 글은 그림 안에, 표 자리 표시는 없다")
    from hwpx import HwpxDocument
    try:
        HwpxDocument.open(str(out_c)); reopened = True
    except Exception:
        reopened = False
    chk(reopened, "D 모양 그대로 HWPX 도 다시 열 수 있다")

    # ── E. 머리말·꼬리말(§1.1) — 머리말·꼬리말 기능으로, 쪽 번호는 필드 ─────────────
    hf = root / "F.pdf"
    d = fitz.open()
    for i in range(4):
        pg_ = d.new_page(width=595, height=842)
        pg_.insert_text((400, 40), f"Asphalt Design Guide p.{i + 11}", fontsize=9, fontname="tiro")   # 글 + 번호
        for k in range(20):
            pg_.insert_text((72, 100 + k * 16), f"Body line {k} on page {i + 1} explains the method.", fontsize=11, fontname="tiro")
        pg_.insert_text((285, 815), f"- {i + 11} -", fontsize=9, fontname="tiro")
    d.save(str(hf)); d.close()
    # Word 편집 가능
    ew = convert.run(str(hf), "docx_edit")
    wd = Document(str(ew)); s0 = wd.sections[0]
    body_t = " ".join(p.text for p in wd.paragraphs)
    chk([p.text for p in s0.header.paragraphs] == ["Asphalt Design Guide p.1"] and "PAGE" in s0.header._element.xml,
        "E Word 머리글에 머리말 글(글 안 번호도 PAGE 필드)", str([p.text for p in s0.header.paragraphs]))
    fx = s0.footer._element.xml
    chk("PAGE" in fx and "- " in "".join(p.text for p in s0.footer.paragraphs), "E Word 바닥글 쪽 번호는 PAGE 필드", fx[-300:])
    chk('w:start="11"' in s0._sectPr.xml, "E 쪽 번호는 원본 번호(11)부터")
    chk("Asphalt Design Guide" not in body_t and "- 12 -" not in body_t, "E 본문에는 머리말·꼬리말 글이 없다")
    chk(abs(s0.header_distance.pt - 30.5) < 2 and abs(s0.footer_distance.pt - 24.5) < 2,
        "E 머리글·바닥글 자리 = 원본 자리", f"{s0.header_distance.pt} {s0.footer_distance.pt}")
    hp = s0.header.paragraphs[0].paragraph_format.left_indent
    chk(hp is not None and abs(hp.pt - (400 - 56.7)) < 3, "E 머리말 가로 자리(들여쓰기)도 원본대로", str(hp.pt if hp else None))
    # Word 모양 그대로(pdf2docx) — 머리말·꼬리말을 지운 사본을 변환하고 머리글·바닥글로
    es = convert.run(str(hf), "docx_shape")
    wd2 = Document(str(es))
    body2 = " ".join(p.text for p in wd2.paragraphs)
    chk([p.text for p in wd2.sections[0].header.paragraphs] == ["Asphalt Design Guide p.1"]
        and "PAGE" in wd2.sections[0].footer._element.xml, "E Word 모양 그대로도 머리글·바닥글(쪽 번호 필드)")
    chk("Asphalt Design Guide" not in body2 and "- 12 -" not in body2, "E pdf2docx 본문에 머리말·꼬리말이 남지 않는다", body2[:100])
    chk(all(s.header.is_linked_to_previous for s in list(wd2.sections)[1:]) and not wd2.sections[0].header.is_linked_to_previous,
        "E 뒤 구역은 첫 구역 머리글을 따른다")
    # HWPX 편집 가능
    eh = convert.run(str(hf), "hwpx_edit")
    x = "".join(section_xml(eh))
    hdr_ = re.search(r"<hp:header .*?</hp:header>", x, re.S)
    ftr_ = re.search(r"<hp:footer .*?</hp:footer>", x, re.S)
    chk(hdr_ and "Asphalt Design Guide" in hdr_.group(0), "E HWPX 머리말 기능에 머리말 글")
    chk(hdr_ and 'numType="PAGE"' in hdr_.group(0) and 'pos="NONE"' in hdr_.group(0),
        "E HWPX 머리말 글 안 번호는 번호 필드(쪽 번호 위치 표시는 끔 — 두 번 그리지 않는다)")
    pnums = re.findall(r'<hp:pageNum [^>]*pos="(?!NONE)[^"]*"[^>]*>', re.sub(r"<hp:(header|footer) .*?</hp:\1>", "", x, flags=re.S))
    chk(len(pnums) == 1 and 'pos="BOTTOM_CENTER"' in pnums[0] and 'sideChar="-"' in pnums[0],
        "E HWPX 번호뿐인 꼬리말(- 12 -)은 한글 '쪽 번호 매기기'(아래 가운데, 줄표)로", str(pnums))
    chk(not ftr_ or "- " not in re.sub(r"<[^>]+>", "", ftr_.group(0)),
        "E HWPX 번호뿐인 꼬리말을 꼬리말 글로 또 넣지 않는다")
    chk('<hp:newNum num="11" numType="PAGE"' in x, "E HWPX 쪽 번호는 원본 번호(11)부터")
    body_x = re.sub(r"<hp:(header|footer) .*?</hp:\1>", "", x, flags=re.S)
    chk("Asphalt Design Guide" not in body_x and "- 12 -" not in body_x, "E HWPX 본문에는 머리말·꼬리말 글이 없다")
    # HWPX 모양 그대로 — 쪽마다 그 쪽 글 그대로, 머리말·꼬리말 영역 = 원본 자리
    esh = convert.run(str(hf), "hwpx_shape")
    xs2 = section_xml(esh)
    chk(len(xs2) == 4 and all(f'<hp:newNum num="{i + 11}" numType="PAGE"' in xs2[i]
                              and re.search(r'<hp:pageNum [^>]*pos="BOTTOM_CENTER"[^>]*sideChar="-"', xs2[i]) for i in range(4)),
        "E HWPX 모양 그대로 — 쪽마다 '쪽 번호 매기기' + 그 쪽 번호로 새 번호")
    chk(all(f"Asphalt Design Guide p.{i + 11}" in xs2[i] for i in range(4)), "E HWPX 모양 그대로 — 쪽마다 그 쪽 머리말 글")
    m0 = re.search(r'<hp:margin [^>]*>', xs2[0]).group(0)
    vals = dict(re.findall(r'(\w+)="(\d+)"', m0))
    chk(int(vals["header"]) > 0 and int(vals["footer"]) > 0 and abs(int(vals["top"]) - 2952) < 200,
        "E HWPX 모양 그대로 — 위 여백·머리말 길이 = 원본 머리말 자리", m0)
    hdr2 = zipfile.ZipFile(str(esh)).read("Contents/header.xml").decode("utf-8")
    pids = re.findall(r'<hh:paraPr [^>]*?\sid="(\d+)"', hdr2)       # \s — snapToGrid="1" 의 id="1" 을 잡지 않게
    chk(len(pids) == len(set(pids)), "E HWPX 모양 그대로 — 문단 서식 번호가 겹치지 않는다(머리말 서식과 본문 서식)",
        str(sorted(p for p in set(pids) if pids.count(p) > 1)))
    b0 = re.sub(r"<hp:(header|footer) .*?</hp:\1>", "", xs2[0], flags=re.S)
    chk("Asphalt Design Guide" not in b0, "E HWPX 모양 그대로 — 본문에는 머리말 글이 없다")
finally:
    shutil.rmtree(root, ignore_errors=True)

print("\n=== " + ("ALL PASS" if not fails else f"FAILURE ({len(fails)})") + " ===")
for m in fails:
    print(" -", m)
sys.stdout.flush()
os._exit(1 if fails else 0)
